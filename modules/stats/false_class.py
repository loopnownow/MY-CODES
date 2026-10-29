"""
false_class.py — 假阳性/假阴性个体分类
═══════════════════════════════════════════════════════════════════
默认：控制台打印宽表 + 写入工程根 false_classification.xlsx（不随 OUT_DIR 走）
两张工作单：
  false_classification — ID | Split | Group | Pattern | 各模型 TP/TN/FP/FN
  exc — 排除名单（默认即本文件；重写时保留已填 ID）
Pattern：all_FN / all_FP（仅当各模型非空单元格全为 FN 或全为 FP）；其余为空
═══════════════════════════════════════════════════════════════════
"""

from __future__ import annotations
import os
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from sklearn.metrics import roc_curve
import statsmodels.api as sm

_POS_CODES = {"TP", "FN"}
_NEG_CODES = {"TN", "FP"}
_VALID_CODES = _POS_CODES | _NEG_CODES
_META_COLS = ("ID", "Split", "Group", "Pattern")


def compute_false_class(
    somescore: np.ndarray,
    label: np.ndarray,
    labels: Optional[Sequence[str]] = None,
    dataset_label: str = "all",
    model_name: str = "Model",
) -> Tuple[pd.DataFrame, Dict]:
    """
    计算个体分类结果，不绘图。

    返回
    ----
    detail_df : 逐样本表 (PatientID, Label, Score, Prob, Threshold, Class, ...)
    summary   : 汇总 dict
    """
    from modules.utils import dry_run as _dry
    print(f"\n═══ FALSE CLASS [{dataset_label} / {model_name}] ═══")
    if _dry.exit_after_plan("false_class", [
        f"dataset={dataset_label} model={model_name}",
        f"n={len(np.asarray(label))}",
        "Youden classify (skipped)",
    ]):
        return pd.DataFrame(), {}
    somescore = np.asarray(somescore, dtype=float)
    y = np.asarray(label).astype(int)
    assert len(somescore) == len(y) and len(somescore) > 0

    X = sm.add_constant(somescore)
    fit = sm.GLM(y, X, family=sm.families.Binomial()).fit()
    prob = np.asarray(fit.predict(X), dtype=float)

    fpr, tpr, thresholds = roc_curve(y, prob)
    youden = tpr - fpr
    idx_y = int(np.argmax(youden))
    thr, sens, spec = float(thresholds[idx_y]), float(tpr[idx_y]), float(1 - fpr[idx_y])
    print(f"  Youden={thr:.4f} (Sens={sens:.3f}, Spec={spec:.3f})")

    pred_class = (prob >= thr).astype(int)
    cls = np.where(
        (y == 1) & (pred_class == 1), "TP",
        np.where(
            (y == 0) & (pred_class == 0), "TN",
            np.where((y == 0) & (pred_class == 1), "FP", "FN"),
        ),
    )
    tp, tn = int((cls == "TP").sum()), int((cls == "TN").sum())
    fp, fn = int((cls == "FP").sum()), int((cls == "FN").sum())
    print(f"  TP={tp} TN={tn} FP={fp} FN={fn}")

    n = len(prob)
    if labels is not None and len(labels) == n:
        pid = [str(x) for x in labels]
    else:
        pid = [str(i + 1) for i in range(n)]

    detail_df = pd.DataFrame({
        "PatientID": pid,
        "Label": y,
        "Score": somescore,
        "Prob": prob,
        "Threshold": thr,
        "Class": cls,
        "Layer": dataset_label,
        "Model": model_name,
    })
    summary = {
        "Layer": dataset_label,
        "Model": model_name,
        "N": n,
        "Youden": thr,
        "Sens": sens,
        "Spec": spec,
        "TP": tp,
        "TN": tn,
        "FP": fp,
        "FN": fn,
    }
    return detail_df, summary


def _row_classes(row: pd.Series, model_cols: Sequence[str]) -> List[str]:
    out: List[str] = []
    for c in model_cols:
        v = row[c]
        if pd.isna(v):
            continue
        s = str(v).strip().upper()
        if s in _VALID_CODES:
            out.append(s)
    return out


def _pattern_for(classes: Sequence[str]) -> str:
    if not classes:
        return ""
    if all(c == "FN" for c in classes):
        return "all_FN"
    if all(c == "FP" for c in classes):
        return "all_FP"
    return ""


def _infer_group_from_classes(
    classes: Sequence[str],
    pos_lab: str,
    neg_lab: str,
) -> Optional[str]:
    if not classes:
        return None
    has_pos = any(c in _POS_CODES for c in classes)
    has_neg = any(c in _NEG_CODES for c in classes)
    if has_pos and has_neg:
        return None
    return pos_lab if has_pos else neg_lab


def _default_display_labels(
    group_labels: Optional[Dict[int, str]] = None,
) -> Tuple[str, str]:
    if group_labels:
        pos = str(group_labels.get(1, "Positive"))
        neg = str(group_labels.get(0, "Negative"))
        return pos, neg
    return "Positive", "Negative"


def enrich_false_classification_wide(
    wide: pd.DataFrame,
    group_labels: Optional[Dict[int, str]] = None,
    label_by_id: Optional[Dict[str, int]] = None,
) -> pd.DataFrame:
    """
    在宽表上追加 Group（第3列）与 Pattern（第4列）。
    Group 优先用 label_by_id + group_labels；否则由 TP/FN vs TN/FP 推断。
    Pattern：all_FN / all_FP / 空。
    """
    if wide is None or len(wide) == 0:
        return wide

    df = wide.copy()
    for c in ("Group", "Pattern"):
        if c in df.columns:
            df = df.drop(columns=[c])

    model_cols = [c for c in df.columns if c not in ("ID", "Split", "PatientID")]
    pos_lab, neg_lab = _default_display_labels(group_labels)

    groups: List[Optional[str]] = []
    patterns: List[str] = []
    for _, row in df.iterrows():
        classes = _row_classes(row, model_cols)
        g: Optional[str] = None
        if label_by_id is not None:
            pid = str(row["ID"]) if "ID" in df.columns else str(row.get("PatientID", ""))
            lab = label_by_id.get(pid)
            if lab is None:
                try:
                    lab = label_by_id.get(int(pid))  # type: ignore[arg-type]
                except Exception:
                    lab = None
            if lab is not None and group_labels:
                g = str(group_labels.get(int(lab), pos_lab if int(lab) == 1 else neg_lab))
            elif lab is not None:
                g = pos_lab if int(lab) == 1 else neg_lab
        if g is None:
            g = _infer_group_from_classes(classes, pos_lab, neg_lab)
        groups.append(g)
        patterns.append(_pattern_for(classes))

    # insert after Split: Group then Pattern
    if "Split" in df.columns:
        split_idx = list(df.columns).index("Split")
        df.insert(split_idx + 1, "Group", groups)
        df.insert(split_idx + 2, "Pattern", patterns)
    else:
        df.insert(1, "Group", groups)
        df.insert(2, "Pattern", patterns)

    front = [c for c in ("ID", "Split", "Group", "Pattern") if c in df.columns]
    rest = [c for c in df.columns if c not in front]
    return df[front + rest]


def build_false_classification_wide(
    details: Dict[str, pd.DataFrame],
    split_by_id: Optional[Dict[str, str]] = None,
    group_labels: Optional[Dict[int, str]] = None,
) -> Optional[pd.DataFrame]:
    """
    组装宽表：ID | Split | Group | Pattern | <model1> | <model2> | ...
    单元格为 TP/TN/FP/FN。仅使用 Layer==all 的结果。
    """
    all_parts = []
    for key, df in details.items():
        if df is None or len(df) == 0:
            continue
        if "Layer" in df.columns:
            sub = df[df["Layer"].astype(str).str.lower() == "all"].copy()
        else:
            if str(key).lower().startswith("train") or str(key).lower().startswith("test"):
                continue
            sub = df.copy()
        if len(sub) == 0:
            continue
        all_parts.append(sub)

    if not all_parts:
        return None

    wide = None
    model_order: List[str] = []
    label_by_id: Dict[str, int] = {}
    for sub in all_parts:
        m = str(sub["Model"].iloc[0])
        if m not in model_order:
            model_order.append(m)
        piece = sub[["PatientID", "Class"]].copy()
        piece = piece.rename(columns={"Class": m})
        piece["PatientID"] = piece["PatientID"].astype(str)
        piece = piece.drop_duplicates(subset=["PatientID"], keep="first")
        if "Label" in sub.columns:
            for pid, lab in zip(
                sub["PatientID"].astype(str),
                sub["Label"].astype(int),
            ):
                if pid not in label_by_id:
                    label_by_id[pid] = int(lab)
        if wide is None:
            wide = piece
        else:
            wide = wide.merge(piece, on="PatientID", how="outer")

    assert wide is not None
    if split_by_id:
        wide["Split"] = wide["PatientID"].map(
            lambda x: split_by_id.get(str(x), split_by_id.get(x, ""))
        )
    else:
        wide["Split"] = ""

    cols = ["PatientID", "Split"] + [m for m in model_order if m in wide.columns]
    for c in wide.columns:
        if c not in cols:
            cols.append(c)
    wide = wide[cols].rename(columns={"PatientID": "ID"})
    return enrich_false_classification_wide(
        wide,
        group_labels=group_labels,
        label_by_id=label_by_id or None,
    )


def print_false_classification(
    details: Dict[str, pd.DataFrame],
    split_by_id: Optional[Dict[str, str]] = None,
    group_labels: Optional[Dict[int, str]] = None,
) -> Optional[pd.DataFrame]:
    """仅打印 false_classification 宽表，不写文件、不输出其他表。"""
    wide = build_false_classification_wide(
        details, split_by_id=split_by_id, group_labels=group_labels,
    )
    if wide is None or len(wide) == 0:
        print("  ⚠ false_classification: no data — skip print")
        return None
    models = [c for c in wide.columns if c not in _META_COLS]
    print(f"\n═══ false_classification (print only)  N={len(wide)}  models={models} ═══")
    with pd.option_context(
        "display.max_rows", None, "display.max_columns", None,
        "display.width", 200, "display.max_colwidth", 40,
    ):
        print(wide.to_string(index=False))
    print("═══ end false_classification ═══\n")
    return wide


# 单元格底色：FN 红、FP 橙；Pattern all_FN / all_FP 同色
_FC_FILLS = {
    "FN": "FFC7CE",       # 浅红
    "FP": "FCE4D6",       # 浅橙
    "all_FN": "FFC7CE",
    "all_FP": "FCE4D6",
    "TP": "C6EFCE",       # 浅绿（可选）
    "TN": "DDEBF7",       # 浅蓝（可选）
}
_FC_FONTS = {
    "FN": "9C0006",
    "FP": "C65911",
    "all_FN": "9C0006",
    "all_FP": "C65911",
    "TP": "006100",
    "TN": "1F4E79",
}


def _apply_false_class_colors(ws) -> int:
    """为 false_classification 工作表的 FN/FP（及 Pattern）上色；返回着色单元格数。"""
    from openpyxl.styles import Font, PatternFill

    fills = {
        k: PatternFill(start_color=v, end_color=v, fill_type="solid")
        for k, v in _FC_FILLS.items()
    }
    fonts = {
        k: Font(color=v, bold=True)
        for k, v in _FC_FONTS.items()
    }
    n_colored = 0
    max_r = ws.max_row or 0
    max_c = ws.max_column or 0
    if max_r < 2 or max_c < 1:
        return 0
    for r in range(2, max_r + 1):
        for c in range(1, max_c + 1):
            cell = ws.cell(r, c)
            val = cell.value
            if val is None:
                continue
            key = str(val).strip()
            if key in fills:
                cell.fill = fills[key]
                cell.font = fonts[key]
                n_colored += 1
    return n_colored


def save_false_classification_xlsx(
    details: Dict[str, pd.DataFrame],
    out_path: str,
    split_by_id: Optional[Dict[str, str]] = None,
    group_labels: Optional[Dict[int, str]] = None,
) -> None:
    """写入 false_classification.xlsx：宽表 + exc 排除单（保留已填 ID）。"""
    from modules.utils import dry_run as _dry
    if _dry.skip_write(out_path, what="false_classification xlsx"):
        return
    wide = build_false_classification_wide(
        details, split_by_id=split_by_id, group_labels=group_labels,
    )
    if wide is None:
        print("  ⚠ false_classification: no all-layer data — skip write")
        return
    models = [c for c in wide.columns if c not in _META_COLS]
    n_all_fn = int((wide["Pattern"] == "all_FN").sum()) if "Pattern" in wide.columns else 0
    n_all_fp = int((wide["Pattern"] == "all_FP").sum()) if "Pattern" in wide.columns else 0
    n_colored = 0
    from modules import config as _cfg
    from modules.stats.data_io import blank_exc_dataframe

    exc_sheet = str(getattr(_cfg, "EXCLUDE_SHEET", "exc") or "exc")
    exc_col = str(getattr(_cfg, "EXCLUDE_COL", "exc") or "exc")
    exc_df = blank_exc_dataframe(exc_col)
    if os.path.isfile(out_path):
        try:
            xl = pd.ExcelFile(out_path)
            names = [str(s) for s in xl.sheet_names]
            hit = exc_sheet if exc_sheet in names else next(
                (s for s in names if s.lower() == exc_sheet.lower()), None
            )
            if hit is not None:
                kept = pd.read_excel(xl, sheet_name=hit)
                if kept is not None and not kept.empty:
                    exc_df = kept
                    print(f"  keep existing exclude sheet {hit!r} ({len(kept)} row(s))")
        except Exception as e:
            print(f"  ⚠ could not reuse existing exc sheet: {e}")

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        wide.to_excel(writer, sheet_name="false_classification", index=False)
        try:
            exc_df.to_excel(writer, sheet_name=exc_sheet, index=False)
        except Exception as e:
            print(f"  ⚠ write exc sheet failed: {e}")
            pd.DataFrame(columns=["exc"]).to_excel(writer, sheet_name="exc", index=False)
        try:
            ws = writer.sheets["false_classification"]
            n_colored = _apply_false_class_colors(ws)
        except Exception as e:
            print(f"  ⚠ false_classification color failed: {e}")
    print(
        f"  ✓ false_classification.xlsx → {out_path}  "
        f"(N={len(wide)}, models={models}, all_FN={n_all_fn}, all_FP={n_all_fp}, "
        f"colored_cells={n_colored})"
    )
