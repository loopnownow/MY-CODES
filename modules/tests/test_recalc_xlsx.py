"""model-validation/recalc_xlsx.py: the scan half, which runs without LibreOffice.

The whole point of the scan is to refuse a false pass. A formula cell openpyxl wrote
has no cached value until a real engine calculates it, so "no error literals" is not
the same as "fine" — the 2026-10-27 QC caught this reporting `success` on a workbook
containing `=1/0`.
"""
from __future__ import annotations

import importlib.util
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
CLI = REPO_ROOT / "model-validation" / "recalc_xlsx.py"


def _load_cli():
    spec = importlib.util.spec_from_file_location("recalc_xlsx_under_test", CLI)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture()
def book(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "T"
    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = "=SUM(A1:A2)"
    wb.save(path)
    return path


def test_uncached_formulas_are_not_reported_as_success(book):
    cli = _load_cli()
    report = cli._scan(book)
    assert report["total_formulas"] == 1
    assert report["total_uncached"] == 1
    assert report["status"] == "uncalculated"


def test_error_literal_is_reported_with_location(tmp_path):
    path = tmp_path / "broken.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["B2"] = "#DIV/0!"
    wb.save(path)

    report = _load_cli()._scan(path)
    assert report["status"] == "errors_found"
    assert report["error_summary"]["#DIV/0!"]["count"] == 1
    assert report["error_summary"]["#DIV/0!"]["locations"] == ["Sheet!B2"]


def test_values_only_workbook_is_success(tmp_path):
    path = tmp_path / "values.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = 2
    wb.save(path)

    report = _load_cli()._scan(path)
    assert report["total_formulas"] == 0
    assert report["status"] == "success"


def test_cli_exit_codes(tmp_path, book, capsys):
    cli = _load_cli()
    # uncalculated
    assert cli.main([str(book), "--check-only", "--json"]) == 3
    # missing file
    assert cli.main([str(tmp_path / "absent.xlsx")]) == 1


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_recalculation_fills_cached_values(book):
    cli = _load_cli()
    assert cli.main([str(book), "--json"]) == 0
    report = cli._scan(book)
    assert report["total_uncached"] == 0


def test_missing_libreoffice_fails_loudly(tmp_path, book, capsys, monkeypatch):
    """No soffice: exit 1 with a message, not a traceback and not exit 0."""
    cli = _load_cli()
    monkeypatch.setattr(cli.shutil, "which", lambda _name: None)
    assert cli.main([str(book)]) == 1
    err = capsys.readouterr().err
    assert "LibreOffice" in err
    assert "--check-only" in err
