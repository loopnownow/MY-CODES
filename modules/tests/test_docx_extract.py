"""docx_extract: text extraction, normalization, and error paths."""
from __future__ import annotations

import os
import sys
from pathlib import Path

import docx
import pytest

# docx_extract: in-repo tools/ first, then ODOC_CODE / GROK_ROOT fallbacks
def _resolve_ms_body() -> Path:
    repo = Path(__file__).resolve().parents[2]
    cands: list[Path] = [repo / "tools" / "write-paper"]  # 2026-11-01 P2: tools/ moved into this repo
    if os.environ.get("ODOC_CODE"):
        cands.append(Path(os.environ["ODOC_CODE"]) / "05-manuscript" / "write-paper")
    if os.environ.get("GROK_ROOT"):
        cands.append(
            Path(os.environ["GROK_ROOT"]) / "0doc" / "_code" / "05-manuscript" / "write-paper"
        )
    for hit in cands:
        if (hit / "docx_extract.py").is_file():
            return hit
    return Path(".")

_MS_BODY = _resolve_ms_body()
if str(_MS_BODY) not in sys.path:
    sys.path.insert(0, str(_MS_BODY))

from docx_extract import _docx_comments, _norm, extract_text


def test_norm_maps_word_controls():
    # \x07 = table cell, \x0b/\x0c = breaks, \r\n = paragraph.
    assert _norm("a\x07b\x0bc\x0cd\r\ne") == "a\tb\nc\nd\ne"


def test_norm_drops_other_controls():
    assert _norm("x\x01y\x02z") == "xyz"


def test_extract_docx_paragraphs_and_tables(tmp_path):
    p = tmp_path / "t.docx"
    d = docx.Document()
    d.add_paragraph("Hello")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "A"
    t.rows[0].cells[1].text = "B"
    d.save(p)

    text = extract_text(str(p))
    assert "Hello" in text
    assert "A\tB" in text


def test_docx_comments_absent_returns_empty(tmp_path):
    p = tmp_path / "n.docx"
    docx.Document().save(p)
    assert _docx_comments(str(p)) == ""


def test_extract_unsupported_extension_raises(tmp_path):
    p = tmp_path / "x.txt"
    p.write_text("hi", encoding="utf-8")
    with pytest.raises(ValueError):
        extract_text(str(p))
