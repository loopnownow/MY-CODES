# -*- coding: utf-8 -*-

r"""Smoke test for the global library and the project settings.



Layout: algorithms live only in the library this checkout ships (`<repo>/modules`).

Projects keep ref/settings.ini under %GROK_ROOT%\<stage>\<project>; there is no

per-project modules/ folder. ZERAD_ROOT / ORAD_ROOT limits the scan to one stage.

"""

from pathlib import Path

import os

import sys



import numpy as np



LIB = Path(os.environ.get("STATS_ROOT") or Path(__file__).resolve().parents[2])  # 仓根（本脚本在 tools/sync/ 下）

if str(LIB) not in sys.path:
    sys.path.insert(0, str(LIB))
from modules.console.paths import grok_root  # GROK_ROOT env -> frozen lab default (C1)
GROK_ROOT = grok_root(LIB / "modules")  # 阶段根父目录；不由 LIB 反推

GOLD = LIB / "modules"

STAGE_NAMES = (

    "drafting", "pre-submit", "in-review", "accepted",


)

_rad = os.environ.get("ZERAD_ROOT") or os.environ.get("ORAD_ROOT")

STAGE_ROOTS = [Path(_rad)] if _rad else [GROK_ROOT / n for n in STAGE_NAMES if (GROK_ROOT / n).is_dir()]

TMP_DIR = Path(os.environ.get("DEL_ROOT") or r"D:\0del") / "harvested" / "smoke"  # 临时输出落 {DEL} 的 harvested 档（0del 根保持只有三档；见 My rule.md §5.3）

# Projects that must carry INPUT and DO_EXCLUDE_SAMPLES in ref/settings.ini.

KEY_PROJECTS = ("ybc_EC", "lxf_RP", "xlm_LG", "lya_CTA")



sys.path.insert(0, str(LIB))

from modules.stats.false_class import compute_false_class, save_false_classification_xlsx  # noqa: E402

from modules.config.ini_io import read_ini  # noqa: E402



failures = []



rng = np.random.RandomState(0)

n = 20

y = rng.binomial(1, 0.5, n)

sc = rng.normal(y, 1, n)

d, _ = compute_false_class(

    sc, y, labels=[f"P{i}" for i in range(n)], dataset_label="all", model_name="M1"

)

TMP_DIR.mkdir(parents=True, exist_ok=True)

tmp = TMP_DIR / "fc_color_test.xlsx"

save_false_classification_xlsx({"M1": d}, str(tmp))



from openpyxl import load_workbook  # noqa: E402



wb = load_workbook(tmp)

ws = wb.active

fills = 0

for row in ws.iter_rows(min_row=2):

    for cell in row:

        if cell.value in ("FN", "FP", "TP", "TN", "all_FN", "all_FP"):

            rgb = None

            try:

                rgb = cell.fill.fgColor.rgb

            except Exception:

                pass

            if rgb and str(rgb) not in ("00000000", "None"):

                fills += 1

print("colored class cells:", fills, "file:", tmp)

if not fills:

    failures.append("no colored FN/FP/TP/TN cells")



data_io = (GOLD / "stats" / "data_io.py").read_text(encoding="utf-8")

for needle in ("apply_sample_exclusion", "SURVIVAL_TIME_COL"):

    ok = needle in data_io

    print("gold data_io", needle, "OK" if ok else "FAIL")

    if not ok:

        failures.append(f"gold data_io lacks {needle}")



n_proj = 0

seen = set()

for root in STAGE_ROOTS:

    for proj in sorted(p for p in root.iterdir() if p.is_dir() and not p.name.startswith((".", "0"))):

        ini = proj / "ref" / "settings.ini"

        if not ini.is_file():

            continue

        n_proj += 1

        seen.add(proj.name)

        data = read_ini(ini)

        md = str(data.get("MODULES_DIR") or "")

        if md and Path(md).resolve() != GOLD.resolve():

            failures.append(f"{root.name}/{proj.name}: MODULES_DIR={md} is not {GOLD}")

        if proj.name in KEY_PROJECTS:

            missing = [k for k in ("INPUT", "DO_EXCLUDE_SAMPLES") if k not in data]

            print(proj.name, "INPUT=", data.get("INPUT"), "missing " + str(missing) if missing else "DO_EXCLUDE_SAMPLES present")

            if missing:

                failures.append(f"{proj.name}: missing {missing}")

print("project settings.ini checked:", n_proj, "MODULES_DIR ->", GOLD)

for name in KEY_PROJECTS:

    if name not in seen:

        failures.append(f"{name}: no ref/settings.ini found")



if failures:

    print("SMOKE FAIL")

    for f in failures:

        print("  ", f)

    sys.exit(1)

print("ALL SMOKE OK")

