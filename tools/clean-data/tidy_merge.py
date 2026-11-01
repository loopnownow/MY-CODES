#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""tidy_merge — 统一「合并」入口，三种策略对应原来的三套合并代码。

策略
----
``keyed_union``    按主键取并集、逐字段冲突解决、输出主表 + 对齐分表 + 冲突日志。
                   原 `lxj_AN/build_merge.py`（profile ``lxj_AN``）。
``source_records`` 主库 + 多源规范化长记录 → 按姓名键回填 / 追加孤儿 / cohort 回推。
                   原 `ssd/clinical_recover.py` + `_sources.py` + `_phases.py`(phase0/1)
                   （profile ``ssd``）。
``modality_outer`` 两个模态表按 Name 外连接（共享临床列优先左表，影像列不重名）。
                   原 `ssd/split_ssd_imputed.py:outer_merge_modalities`。

用法::

    python tidy_merge.py --profile lxj_AN
    python tidy_merge.py --profile ssd
    python tidy_merge.py --strategy modality_outer --left A.xlsx --right B.xlsx --out C.xlsx

``keyed_union`` / ``source_records`` 都不修改任何原始文件。
"""
from __future__ import annotations

import argparse
import collections
import sys
import warnings
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent))

from tidy_io import (  # noqa: E402
    KEY_FILL,
    is_blank,
    load_keyed,
    mark_key_column,
    name_key_compact,
    norm,
    norm_name,
    normalize_sex,
    open_book,
    read_frame,
    set_if,
    style_header,
    to_num,
    values_conflict,
    write_frames,
)
from tidy_profiles import (  # noqa: E402
    CANON_FIELDS,
    MERGE_VALUE_FIELDS,
    NUMERIC_FIELDS,
    LxjAnProfile,
    SourceSpec,
    SsdProfile,
    get_profile,
)

__all__ = [
    "merge_keyed_union",
    "merge_source_records",
    "outer_merge_modalities",
    "build_source_records",
    "phase0_inventory",
    "phase1_real_merge",
    "main",
]


# ═════════════════════════════════════════════════════════════════
# 策略 1：keyed_union（lxj_AN）
# ═════════════════════════════════════════════════════════════════
def merge_keyed_union(
    profile: LxjAnProfile,
    out_path: Optional[Path] = None,
    verbose: bool = True,
) -> Tuple[Path, List[str]]:
    """按主键并集合并多源 .xls → 单工作簿（主表 + 对齐分表 + 冲突与来源日志）。

    逐行等价于 `lxj_AN/build_merge.py`。
    """
    p = profile
    log: List[str] = []

    def L(msg: str = "") -> None:
        log.append(msg)
        if verbose:
            print(msg)

    out_path = Path(out_path) if out_path else p.root / p.out_merge

    # ── 打开三本工作簿 ──
    bkA = open_book(p.path(p.file_a))
    bkB = open_book(p.path(p.file_b))
    bkC = open_book(p.path(p.file_c))
    shB = bkB.sheet(0)
    shC = bkC.sheet(0)
    b_hdr_full = bkB.header(shB, 0)
    c_hdr_full = bkC.header(shC, 0)

    # ── B / C（单表源） ──
    hB, rB, dB, oB = load_keyed(bkB, shB, 0, skip_keys=p.skip_keys)
    L(f"[B] 读入 {len(rB)} 个编号（已剔除测试记录 就诊号=1/姓名=测试）；"
      f"重复编号 {len(dB)} 个，已去重 {sum(d - 1 for _, _, d, _ in dB)} 行")
    for x in dB:
        L(f"    B 重复: 就诊号={x[1]} 出现{x[2]}次 ({x[3]})")

    hC, rC, dC, oC = load_keyed(bkC, shC, 1, skip_keys=p.skip_keys)
    L(f"[C] 读入 {len(rC)} 个编号（已剔除测试记录 就诊号=1/姓名=测试）；重复编号 {len(dC)} 个")

    # ── A（多表源） ──
    A_data: Dict[str, Any] = {}
    for spec in p.a_sheets:
        sh = bkA.sheet(spec.sheet)
        hdr, rows, dups, _order = load_keyed(bkA, sh, spec.key_col)
        A_data[spec.sheet] = (hdr, rows)
        L(f"[A:{spec.sheet}] 读入 {len(rows)} 个编号；重复编号 {len(dups)} 个，"
          f"已去重 {sum(d - 1 for _, _, d, _ in dups)} 行")
        for x in dups:
            L(f"    A/{spec.sheet} 重复: 体检编号={x[1]} 出现{x[2]}次 ({x[3]})")

    # ── 行基准：全部唯一编号的并集 ──
    allA: set = set()
    for sname in A_data:
        allA |= set(A_data[sname][1])
    core = set(rB) | set(rC)
    ids = sorted(allA | core)
    L("")
    L(f"[行基准] A独立编号={len(allA)}  B|C编号={len(core)}  "
      f"交集={len(allA & core)}  并集(总表行数)={len(ids)}")
    L(f"   仅A有(无眼动数据)={len(allA - core)}    "
      f"仅B/C有(无碳纤维数据)={len(core - allA)}")

    # ── 组别票决 ──
    def resolve_group(key: str) -> Tuple[str, List[Tuple[str, str]]]:
        votes: List[Tuple[str, str]] = []
        for spec in p.a_sheets:
            if spec.group_col is None:
                continue
            row = A_data[spec.sheet][1].get(key)
            if not row:
                continue
            g = norm(row[spec.group_col]) if spec.group_col < len(row) else ""
            if g:
                votes.append((spec.sheet, g))
        if not votes:
            return "", votes
        cnt = collections.Counter(v for _, v in votes)
        return cnt.most_common(1)[0][0], votes

    group_conflicts: List[Tuple[str, str, str]] = []
    age_diff: List[Tuple[str, str, str]] = []
    for k in ids:
        g, votes = resolve_group(k)
        if len({v for _, v in votes}) > 1:
            group_conflicts.append((k, "; ".join(f"{s}={v}" for s, v in votes), g))
        # 年龄对比：A 表 vs B/C
        a_age = ""
        for sname in p.a_age_sheets:
            hdr, rows = A_data[sname]
            row = rows.get(k)
            if row and len(row) > 3 and row[3] != "":
                a_age = norm(row[3])
                break
        bc_age = (
            norm(rB[k][3]) if k in rB and len(rB[k]) > 3 and rB[k][3] != ""
            else (norm(rC[k][4]) if k in rC and len(rC[k]) > 4 and rC[k][4] != "" else "")
        )
        if a_age and bc_age and str(a_age) != str(bc_age):
            age_diff.append((k, bc_age, a_age))

    L(f"[组别冲突] {len(group_conflicts)} 个编号在A各表间\"实验组/对照组\"不一致")
    L(f"[年龄差异] {len(age_diff)} 个编号 A表年龄 与 B/C 不同（已按B/C为准）")

    # ── 字段级冲突（A 三表间） ──
    def _a_field(key: str, colname: str) -> Dict[str, str]:
        vals: Dict[str, str] = {}
        for sname in p.a_field_sheets:
            hdr, rows = A_data[sname]
            rw = rows.get(key)
            if rw and colname in hdr:
                v = norm(rw[hdr.index(colname)])
                if v:
                    vals[sname] = v
        return vals

    field_conf: List[Tuple[str, str, Dict[str, str]]] = []
    for field, hname in p.field_names:
        for k in ids:
            vals = _a_field(k, hname)
            if len(set(vals.values())) > 1:
                field_conf.append((k, field, vals))
                moca = vals.get("MoCA详细")
                swapped = (
                    (field == "文化程度" and moca in ("是", "否"))
                    or (field == "是否为右利手" and moca in
                        ("小学", "初中", "高中", "中专", "大专", "专科", "本科", "硕士", "博士"))
                )
                if swapped:
                    L(f"    !! {k} 在 MoCA详细 表中「文化程度/是否为右利手」两列取值疑似错位: {vals}")
    L(f"[字段冲突] 文化程度/是否为右利手 在A三表间不一致 {len(field_conf)} 处"
      f"（总表取 认知量表得分>MMSE详细>MoCA详细 的首个非空值）")

    # ── 主表列结构 ──
    b_extra_idx = list(range(4, len(b_hdr_full)))          # 0就诊号 1姓名 2性别 3年龄 已提升
    c_extra_idx = [0] + list(range(7, 17))                 # C_ID + 诊断..创建时间
    master_cols: List[str] = list(p.head_cols)
    for c in b_extra_idx:
        master_cols.append(b_hdr_full[c])
    master_cols.append("C_系统内码")
    for c in c_extra_idx:
        if c == 0:
            continue
        master_cols.append("C_" + c_hdr_full[c])
    for spec in p.a_sheets:
        hdr = A_data[spec.sheet][0]
        for c, nm in spec.extras:
            master_cols.append(spec.prefix + (nm if nm else hdr[c]))

    seen = collections.Counter()
    deduped: List[str] = []
    for c in master_cols:
        seen[c] += 1
        deduped.append(c if seen[c] == 1 else f"{c}_{seen[c]}")
    master_cols = deduped
    L(f"[总表列数] {len(master_cols)}")

    # ── 逐行装配 ──
    def getA(key: str, sname: str, col: int) -> Any:
        hdr, rows = A_data[sname]
        row = rows.get(key)
        if not row or col >= len(row):
            return ""
        return row[col]

    srcflag: collections.Counter = collections.Counter()
    master_rows: List[List[Any]] = []
    for k in ids:
        b = rB.get(k)
        c = rC.get(k)
        name = (norm(b[1]) if b and len(b) > 1 and b[1] != ""
                else (norm(c[2]) if c and len(c) > 2 else ""))
        sex = (norm(b[2]) if b and len(b) > 2 and b[2] != ""
               else (norm(c[3]) if c and len(c) > 3 and c[3] != "" else ""))
        age = (norm(b[3]) if b and len(b) > 3 and b[3] != ""
               else (norm(c[4]) if c and len(c) > 4 and c[4] != "" else ""))
        if not sex or not age:
            for sname in p.a_demo_sheets:
                hdr, rows = A_data[sname]
                rw = rows.get(k)
                if not rw:
                    continue
                ci = hdr.index("性别") if "性别" in hdr else None
                ai = hdr.index("年龄") if "年龄" in hdr else None
                if not sex and ci is not None and ci < len(rw):
                    sex = norm(rw[ci])
                if not age and ai is not None and ai < len(rw):
                    age = norm(rw[ai])

        a_age = ""
        for sname in p.a_age_sheets:
            hdr, rows = A_data[sname]
            rw = rows.get(k)
            if rw and "年龄" in hdr:
                v = norm(rw[hdr.index("年龄")])
                if v:
                    a_age = v
                    break
        edu_c = norm(c[6]) if c and len(c) > 6 else ""
        cult = ""
        for sname in p.a_field_sheets:
            if sname not in A_data:
                continue
            hdr, rows = A_data[sname]
            rw = rows.get(k)
            if rw and "文化程度" in hdr:
                v = norm(rw[hdr.index("文化程度")])
                if v:
                    cult = v
                    break
        hand = ""
        for sname in p.a_field_sheets:
            if sname not in A_data:
                continue
            hdr, rows = A_data[sname]
            rw = rows.get(k)
            if rw and "是否为右利手？" in hdr:
                v = norm(rw[hdr.index("是否为右利手？")])
                if v:
                    hand = v
                    break

        inA, inB, inC = k in allA, k in rB, k in rC
        src = "+".join([x for x, f in (("A", inA), ("B", inB), ("C", inC)) if f])
        srcflag[src] += 1

        vals = {
            "编号": k, "组别": resolve_group(k)[0], "姓名": name, "性别": sex,
            "年龄": age, "年龄_A表": a_age,
            "出生日期": norm(c[5]) if c and len(c) > 5 else "",
            "教育程度": edu_c, "文化程度": cult, "是否为右利手": hand, "数据来源": src,
        }
        row: List[Any] = []
        for col in master_cols:
            if col in vals:
                row.append(vals[col])
                continue
            v: Any = ""
            if b is not None and col in b_hdr_full and b_hdr_full.index(col) in b_extra_idx:
                v = norm(b[b_hdr_full.index(col)])
            elif col.startswith("C_"):
                base = col[2:]
                if base == "系统内码":
                    v = norm(c[0]) if c else ""
                elif c is not None and base in c_hdr_full:
                    v = norm(c[c_hdr_full.index(base)])
            else:
                for spec in p.a_sheets:
                    hdr = A_data[spec.sheet][0]
                    for ci, nm in spec.extras:
                        outn = spec.prefix + (nm if nm else hdr[ci])
                        if outn == col:
                            v = norm(getA(k, spec.sheet, ci))
                            break
                    if v != "":
                        break
            row.append(v)
        master_rows.append(row)

    L(f"[数据来源分布] {dict(srcflag)}")

    # ── 列填充审计 ──
    if len(set(master_cols)) != len(master_cols):
        L("!! 警告: 总表存在重名列")
    empties = [
        col for i, col in enumerate(master_cols)
        if sum(1 for r in master_rows if str(r[i]).strip() != "") == 0
    ]
    L(f"[空列] {len(empties)} 列在总表中全空: {empties}" if empties
      else "[空列] 无（每列都有数据）")

    # ── 写出 ──
    wb = Workbook()
    ws = wb.active
    ws.title = p.master_sheet
    ws.append(master_cols)
    for r in master_rows:
        ws.append(r)
    style_header(ws, len(master_cols))
    ws.column_dimensions["A"].width = 15
    for col in ("B", "C", "D", "E", "F", "K"):
        ws.column_dimensions[col].width = 12
    mark_key_column(ws, 1, width=15)

    # 对齐分表
    def align_sheet(title: str, hdr: Sequence[str], rows: Dict[str, List[Any]]) -> None:
        w = wb.create_sheet(title)
        src_cols = list(hdr)
        w.append(["编号"] + src_cols)
        for k in ids:
            rw = rows.get(k)
            w.append([k] + ([norm(x) for x in rw] if rw else [""] * len(src_cols)))
        style_header(w, len(src_cols) + 1)
        w.column_dimensions["A"].width = 15
        w.column_dimensions["A"].fill = KEY_FILL

    align_sheet("B_测试数据记录", b_hdr_full, {k: v for k, v in rB.items()})
    align_sheet("C_患者列表", c_hdr_full, {k: v for k, v in rC.items()})
    for spec in p.a_sheets:
        hdr, rows = A_data[spec.sheet]
        align_sheet("A_" + spec.sheet, hdr, rows)

    # 冲突 / 来源日志
    ws = wb.create_sheet(p.log_sheet)
    ws.append(["项目", "说明"])

    def add(k: Any, v: Any) -> None:
        ws.append([k, v])

    add("总表行数", f"{len(ids)} 行（A、B、C 三文件全部唯一编号的并集）")
    add("行基准", "按“编号”对齐，非按行号；A用体检编号，B/C用就诊号，两者同一编号体系")
    add("", "")
    add("—— 来源分布 ——", "")
    for k, v in sorted(srcflag.items(), key=lambda x: -x[1]):
        add(f"{k}", f"{v} 人")
    add("", "")
    add("—— A与B/C的匹配情况 ——", "")
    add("A 独立编号数", len(allA))
    add("B/C 编号数", len(core))
    add("交集", len(allA & core))
    add("仅A有（未做眼动）", f"{len(allA - core)} 人，B/C相关列全空")
    add("仅B/C有（无碳纤维体检）", f"{len(core - allA)} 人，A相关列全空")
    add("", "")
    add("—— 去重处理 ——", "")
    add("规则", "同编号多行时保留非空字段最多的一行；完全相同者去重")
    for x in dB:
        add(f"B 重复 就诊号={x[1]}", f"出现{x[2]}次，{x[3]}，已保留最完整一行")
    for spec in p.a_sheets:
        sh_ = bkA.sheet(spec.sheet)
        _, _, dups, _ = load_keyed(bkA, sh_, spec.key_col)
        for x in dups:
            add(f"{spec.sheet} 重复 编号={x[1]}", f"出现{x[2]}次，{x[3]}，已保留最完整一行")
    add("剔除", "B表 就诊号=1 / 姓名=测试 的测试记录（该记录在C表中同样存在，已一并剔除）")
    add("注意", "C表 编号=11（余婷婷，女，36岁）为异常短编号，仅存在于C表，已保留，请确认是否为正式受试者")
    add("", "")
    add("—— 组别冲突 ——", "")
    add("冲突数", f"{len(group_conflicts)} 个（取多数票；相同票数取A表首次出现值）")
    for k, detail, win in group_conflicts:
        add(f"编号 {k}", f"{detail} → 采用「{win}」")
    add("", "")
    add("—— 年龄差异 ——", "")
    add("说明", f"{len(age_diff)} 个编号 A表年龄 与 B/C 不一致；"
               f"总表\"年龄\"以B/C为准，\"年龄_A表\"保留A原值")
    for k, bc, aa in age_diff:
        add(f"编号 {k}", f"B/C={bc}，A表={aa}")
    add("", "")
    add("—— 字段取值冲突（文化程度 / 是否为右利手）——", "")
    add("说明", f"{len(field_conf)} 处不一致；总表按「认知量表得分 → MMSE详细 → MoCA详细」取首个非空值")
    for k, field, vals in field_conf:
        detail = "；".join(f"{sn}={v}" for sn, v in vals.items())
        flag = "  ★疑似该两列取值错位" if (
            field == "文化程度" and vals.get("MoCA详细") in ("是", "否")) else ""
        add(f"{field} 编号 {k}", f"{detail}{flag}")
    add("", "")
    add("—— 字段口径提示 ——", "")
    add("性别/年龄", "B与C完全一致，优先取B；缺失时依次回退C、A表")
    add("教育程度 vs 文化程度", "C表用\"专科/其他\"，A表用\"大专/中专\"，编码未统一，两列分别保留原值")
    add("C表空列", "诊断/病史/民族/视力/矫正后视力/身高/体重/利手/睡眠时长 在C表100%为空，已在总表以 C_ 前缀保留")
    add("A表空列", "MMSE_记忆力-3分 与 MoCA_记忆-不计分 在A表中原本就无任何数据（MoCA该项标注\"不计分\"）")
    add("B表列", "眼动范式参数与指标，共49列，列名保持原样")
    style_header(ws, 2)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 95
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=False, vertical="top")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    L("")
    L(f"SAVED {out_path}")
    L(f"sheets: {wb.sheetnames}")
    L(f"master: {len(master_cols)} cols x {len(master_rows)} rows")
    return out_path, log


# ═════════════════════════════════════════════════════════════════
# 策略 2：source_records（SSD）
# ═════════════════════════════════════════════════════════════════
def build_source_records(specs: Sequence[SourceSpec], profile: SsdProfile,
                         verbose: bool = True) -> pd.DataFrame:
    """把多个源表规范化成 canonical 长表（原 ``load_all_sources``）。

    行字典只预置 ``CANON_FIELDS``；`height_cm` / `weight_kg` / `followup` 等
    ``CANON_EXTRA`` 字段**按需出现**——旧实现用 ``pd.DataFrame(rows)`` 建表，
    某个额外列若所有源都没给值就根本不会出现在 `source_long.csv` 里，此处保持一致。
    """
    cols = list(CANON_FIELDS)
    rows: List[Dict[str, Any]] = []
    for spec in specs:
        path = Path(spec.path)
        if not path.is_absolute():
            path = profile.rel(spec.path)
        df = read_frame(path, sheet=spec.sheet, required=spec.required)
        if df is None:
            if verbose:
                print(f"  [skip] {spec.source}: {path} 不存在")
            continue
        if verbose:
            print(f"  [src] {spec.source}: {path.name}::{spec.sheet} n={len(df)}")
        for _, r in df.iterrows():
            rec: Dict[str, Any] = {k: np.nan for k in cols}
            rec["source"] = spec.source
            if spec.cohort is not None:
                rec["cohort_src"] = spec.cohort
            elif spec.cohort_col is not None:
                rec["cohort_src"] = r.get(spec.cohort_col)
            if spec.name_en:
                set_if(rec, "name_en", r.get(spec.name_en))
            if spec.name_cn:
                set_if(rec, "name_cn", r.get(spec.name_cn))
            if spec.subject_id:
                set_if(rec, "subject_id", r.get(spec.subject_id))
            if spec.sex:
                sex = normalize_sex(r.get(spec.sex))
                if sex:
                    rec["sex"] = sex
            for canon, col in spec.numeric:
                set_if(rec, canon, r.get(col), numeric=True)
            for canon, col in spec.text:
                set_if(rec, canon, r.get(col))
            rows.append(rec)
    src = pd.DataFrame(rows)
    src["_name_key"] = src["name_en"].map(name_key_compact)
    src["_name_norm"] = src["name_en"].map(norm_name)
    return src


def phase0_inventory(main: pd.DataFrame, src: pd.DataFrame,
                     profile: SsdProfile) -> str:
    """盘点报告（原 ``phase0_inventory``）。"""
    out = profile.out
    lines: List[str] = []
    lines.append("# SSD 临床数据盘点报告 (00_data_inventory)\n")
    lines.append(f"- 生成目录: `{out}`\n")
    lines.append(f"- 随机种子（后续合成）: `{profile.seed}`\n")

    lines.append("\n## 1. 材料清单\n")
    lines.append("| 文件 | 说明 |\n|------|------|\n")
    lines.append(f"| `{profile.main_file}` | 主临床库 shape={main.shape} |\n")
    lines.append("| `0del/0list.xlsx` | SSD/HC/wen_clin 名单 |\n")
    lines.append("| `0del/SSD患者数据合并结果.xlsx` | dyf/wxt/HC 中文源 |\n")
    lines.append("| `SSD.xlsx` | 多模态 ROI + clinical sheet（本次只登记临床） |\n")
    lines.append("| 两篇论文 docx | 临床靶表 mean±SD |\n")

    lines.append("\n## 2. 主库 cohort 分布\n")
    vc = main[profile.main_cohort_col].value_counts(dropna=False)
    lines.append("| cohort | n |\n|--------|---|\n")
    for k, v in vc.items():
        lines.append(f"| {k} | {v} |\n")

    lines.append("\n## 3. 主库字段非缺失率\n")
    lines.append("| 字段 | 全体 | SSD | HC | unknown |\n|------|------|-----|----|----------|\n")

    def rate(sub: pd.DataFrame, c: str) -> str:
        if len(sub) == 0:
            return "—"
        nn = sub[c].notna().sum()
        return f"{nn}/{len(sub)} ({100 * nn / len(sub):.0f}%)"

    for c in main.columns:
        lines.append(
            f"| `{c}` | {rate(main, c)} | "
            f"{rate(main[main[profile.main_cohort_col] == 'SSD'], c)} | "
            f"{rate(main[main[profile.main_cohort_col] == 'HC'], c)} | "
            f"{rate(main[main[profile.main_cohort_col] == 'unknown'], c)} |\n"
        )

    lines.append("\n## 4. 多源记录数（规范化后）\n")
    lines.append("| source | n_rows | n_named |\n|--------|--------|----------|\n")
    for s, g in src.groupby("source"):
        lines.append(f"| {s} | {len(g)} | {g['_name_key'].notna().sum()} |\n")

    lines.append("\n## 5. 与论文临床靶表对照\n")
    lines.append("### 论文1（应用研究 表1）\n")
    lines.append("- SSD n=101: age 47.30±13.60, edu 11.56±3.29, PHQ-15 8.60±3.45, SSS 33.52±6.05, M/F 29/72\n")
    lines.append("- HC n=36: age 44.19±12.22, edu 13.14±2.94, PHQ-15 1.81±1.90, SSS 22.28±2.56, M/F 18/18\n")
    lines.append("\n### 论文2（毕业论文 part1 表1-1）\n")
    lines.append("- SSD n=56: age 47.63±12.17, M/F 25/31, edu 10.82±3.37, MMSE 27.25±1.27, "
                 "PHQ 8.32±2.56, SSS 35.48±3.02, GAD-7 5.36±3.98, HAMD 6.66±2.98\n")
    lines.append("- HC n=46: age 42.74±14.69, M/F 23/23, edu 11.04±3.85, MMSE 27.7±1.11, "
                 "PHQ 1.35±0.73, SSS 22.26±0.74, GAD-7 3.98±1.56, HAMD 4.13±1.75\n")

    lines.append("\n### 主库现测（非缺失子集）\n")
    for gname in ["SSD", "HC"]:
        sub = main[main[profile.main_cohort_col] == gname]
        lines.append(f"\n**{gname}** n={len(sub)}\n\n")
        lines.append("| 变量 | n | mean±SD |\n|------|---|----------|\n")
        for c in ["age", "education_years", "mmse", "phq15", "sss", "gad7", "sds", "sas", "bmi"]:
            if c not in sub.columns:
                continue
            v = pd.to_numeric(sub[c], errors="coerce").dropna()
            lines.append(f"| {c} | 0 | — |\n" if len(v) == 0
                         else f"| {c} | {len(v)} | {v.mean():.2f}±{v.std():.2f} |\n")

    lines.append("\n## 6. 影像 ROI 登记（不补全）\n")
    ssd_xlsx = profile.rel("SSD.xlsx")
    if ssd_xlsx.exists():
        for sheet in list(profile.split_sheets) + ["clinical"]:
            try:
                df = pd.read_excel(ssd_xlsx, sheet_name=sheet)
                g = df["Group"].value_counts().to_dict() if "Group" in df.columns else {}
                lines.append(f"- `{sheet}`: n={len(df)}, Group={g}\n")
            except Exception as e:  # noqa: BLE001  （与原实现一致：缺表也写一行 error）
                lines.append(f"- `{sheet}`: error {e}\n")
    lines.append("\n> 影像为另一套 ROI 残片；rs-fMRI/VBM 原始数据已丢失且无法从论文还原。\n")

    lines.append("\n## 7. 可合成 vs 不可合成\n")
    lines.append("**可合成（仅补缺/演示）**: " + profile.synth_note + "\n\n")
    lines.append("**永不合成**: 姓名、电话、地址、主诉原文、用药、既往史等隐私/自由文本\n")

    lines.append("\n## 8. 关键结论\n")
    lines.append("1. 主库已基本吸收各源中的量表数值；再合并主要收益在 cohort 回推、零星字段与冲突标注。\n")
    lines.append("2. HAMD 在真实源中完全缺失；仅能合成。\n")
    lines.append("3. 主库 N（SSD138/HC69）> 论文最终分析 N；发表样本是质控子集。\n")
    lines.append("4. unknown=42 需规则回推，不能全部当缺失丢弃。\n")
    return "".join(lines)


def phase1_real_merge(main: pd.DataFrame, src: pd.DataFrame,
                       profile: SsdProfile) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """主库 + 多源真实值合并（原 ``phase1_real_merge``）。返回 (merged, conflicts)。"""
    base = main.copy()
    base["_name_key"] = base["name_en"].map(name_key_compact)
    base["name_cn"] = np.nan
    base["hamd"] = np.nan
    base["cohort_final"] = base["cohort"].astype(str)
    base["cohort_rule"] = "main_cohort"
    base["n_sources"] = 1
    base["sources_merged"] = base.get(
        "data_sources", pd.Series([""] * len(base))).astype(str)

    src_named = src[src["_name_key"].notna()].copy()
    conflicts: List[Dict[str, Any]] = []
    fill_counts = {f: 0 for f in MERGE_VALUE_FIELDS}
    groups = {k: g for k, g in src_named.groupby("_name_key")}

    for i, row in base.iterrows():
        key = row["_name_key"]
        if not key or key not in groups:
            continue
        g = groups[key]
        sources = sorted(g["source"].dropna().unique().tolist())
        extra = [s for s in sources if s != "clinical_merged_en"]
        base.at[i, "n_sources"] = 1 + len(extra)
        old_src = str(base.at[i, "sources_merged"] or "")
        merged_src = ";".join([x for x in [old_src] + extra if x and x != "nan"])
        base.at[i, "sources_merged"] = merged_src

        for field in MERGE_VALUE_FIELDS:
            if field not in g.columns:
                continue
            cur = base.at[i, field] if field in base.columns else np.nan
            # 原实现：pd.isna(cur) or 空白串。注意**不能**用 is_blank 代替——
            # 日期列（datetime64）缺失是 NaT，pd.isna 认、is_blank 不认。
            cur_missing = pd.isna(cur) or (isinstance(cur, str) and not str(cur).strip())

            cands: List[Tuple[str, Any]] = []
            for _, sr in g.iterrows():
                if sr["source"] == "clinical_merged_en":
                    continue
                val = sr.get(field)
                if field == "sex":
                    if not isinstance(val, str) or val not in ("male", "female"):
                        val = normalize_sex(val)
                    if val is None:
                        continue
                elif field in NUMERIC_FIELDS:
                    val = to_num(val)
                    if np.isnan(val):
                        continue
                else:
                    if is_blank(val):
                        continue
                cands.append((sr["source"], val))

            if not cands:
                continue

            if cur_missing:
                vals = [c[1] for c in cands]
                chosen = vals[0]
                for s1, v1 in cands[1:]:
                    if values_conflict(chosen, v1):
                        conflicts.append({
                            "name_en": row["name_en"], "field": field,
                            "main_value": cur, "source_a": cands[0][0], "value_a": chosen,
                            "source_b": s1, "value_b": v1,
                            "resolution": "fill_from_first_source",
                        })
                base.at[i, field] = chosen
                fill_counts[field] = fill_counts.get(field, 0) + 1
            else:
                for s1, v1 in cands:
                    if values_conflict(cur, v1):
                        conflicts.append({
                            "name_en": row["name_en"], "field": field,
                            "main_value": cur, "source_a": "main", "value_a": cur,
                            "source_b": s1, "value_b": v1, "resolution": "keep_main",
                        })

    # ── 主库没有、但源里有姓名与可用数据的孤儿 ──
    main_keys = set(base["_name_key"].dropna())
    orphan_rows: List[Dict[str, Any]] = []
    for key, g in groups.items():
        if key in main_keys:
            continue
        g2 = g[g["source"] != "clinical_merged_en"]
        if g2.empty:
            continue
        new: Dict[str, Any] = {c: np.nan for c in base.columns}
        new["name_en"] = g2.iloc[0].get("name_en")
        new["_name_key"] = key
        new["cohort"] = "unknown"
        new["cohort_final"] = "unknown"
        new["cohort_rule"] = "orphan_source"
        new["n_sources"] = g2["source"].nunique()
        new["sources_merged"] = ";".join(sorted(g2["source"].dropna().unique()))
        new["data_sources"] = new["sources_merged"]
        for field in MERGE_VALUE_FIELDS:
            for _, sr in g2.iterrows():
                val = sr.get(field)
                if field == "sex":
                    val = val if val in ("male", "female") else normalize_sex(val)
                    if val is None:
                        continue
                elif field in NUMERIC_FIELDS:
                    val = to_num(val)
                    if np.isnan(val):
                        continue
                else:
                    if is_blank(val):
                        continue
                # 原实现用 pd.isna（NaT 也算空），不用 is_blank
                if pd.isna(new.get(field)):
                    new[field] = val
        for _, sr in g2.iterrows():
            # 与原实现一致：这里用 pd.notna（空白串也算"有值"），不是 is_blank
            if pd.notna(sr.get("subject_id")):
                new["subject_id"] = sr.get("subject_id")
                break
            if pd.notna(sr.get("name_cn")):
                new["name_cn"] = sr.get("name_cn")
        has_data = any(
            not is_blank(new.get(f)) for f in ["age", "phq15", "sss", "mmse", "gad7", "sex"]
        )
        if has_data:
            orphan_rows.append(new)

    if orphan_rows:
        base = pd.concat([base, pd.DataFrame(orphan_rows)], ignore_index=True)

    # ── cohort 回推 ──
    valid = set(profile.valid_cohorts)
    for i, row in base.iterrows():
        cur = str(row.get("cohort_final", row.get("cohort", "unknown")))
        if cur in valid:
            base.at[i, "cohort_final"] = cur
            if base.at[i, "cohort_rule"] == "main_cohort" or pd.isna(base.at[i, "cohort_rule"]):
                base.at[i, "cohort_rule"] = "main_or_existing"
            continue

        key = row.get("_name_key")
        votes: List[str] = []
        if key and key in groups:
            votes = (groups[key]["cohort_src"].dropna().astype(str).str.upper().tolist())
            votes = [v for v in votes if v in valid]

        if votes:
            ssd_n = votes.count("SSD")
            hc_n = votes.count("HC")
            if ssd_n > hc_n:
                base.at[i, "cohort_final"] = "SSD"
                base.at[i, "cohort_rule"] = f"source_vote_SSD_{ssd_n}_vs_HC_{hc_n}"
                continue
            if hc_n > ssd_n:
                base.at[i, "cohort_final"] = "HC"
                base.at[i, "cohort_rule"] = f"source_vote_HC_{hc_n}_vs_SSD_{ssd_n}"
                continue
            if ssd_n == hc_n and ssd_n > 0:
                base.at[i, "cohort_final"] = "unknown"
                base.at[i, "cohort_rule"] = f"source_vote_tie_{ssd_n}"
                continue

        phq = to_num(row.get("phq15"))
        sss = to_num(row.get("sss"))
        if not np.isnan(phq) and phq >= profile.infer_ssd_phq_min:
            base.at[i, "cohort_final"] = "SSD"
            base.at[i, "cohort_rule"] = "inferred_phq15_ge_5"
            continue
        if (not np.isnan(phq) and not np.isnan(sss)
                and phq <= profile.infer_hc_phq_max and sss <= profile.infer_hc_sss_max):
            base.at[i, "cohort_final"] = "HC"
            base.at[i, "cohort_rule"] = "inferred_low_phq_sss"
            continue

        base.at[i, "cohort_final"] = "unknown"
        base.at[i, "cohort_rule"] = "unresolved"

    conflict_df = pd.DataFrame(conflicts)
    base["is_real_row"] = 1
    base.attrs["meta"] = {
        "fill_counts": fill_counts,
        "n_orphans": len(orphan_rows),
        "n_conflicts": len(conflicts),
    }
    return base, conflict_df


def merge_source_records(profile: SsdProfile, out_dir: Optional[Path] = None,
                         verbose: bool = True) -> Dict[str, Any]:
    """主库 + 多源真实合并，写盘点报告 / 长表 / 真值合并表。

    等价于原 `ssd/clinical_recover.py` 的 Phase 0 + Phase 1。
    """
    out = Path(out_dir) if out_dir else profile.out
    # 先校验输入再建目录：否则"主库缺失"会被"输出目录不可写"的错误盖掉（旧脚本即如此）
    if not profile.main_path.exists():
        raise FileNotFoundError(
            f"主库不存在: {profile.main_path}\n"
            f"  请把 SSD 子工程根目录用环境变量 {profile.root_profile.upper()}_ROOT 指过去"
            f"（当前={profile.root}）")
    out.mkdir(parents=True, exist_ok=True)

    if verbose:
        print("Loading main + sources...")
    main = read_frame(profile.main_path, sheet=0)
    specs = (profile.main_spec,) + tuple(profile.source_specs)
    src = build_source_records(specs, profile, verbose=verbose)
    src.to_csv(out / profile.out_source_long, index=False, encoding="utf-8-sig")
    if verbose:
        print(f"  main={main.shape} source_long={src.shape}")

    if verbose:
        print("Phase 0 inventory...")
    (out / profile.out_inventory).write_text(phase0_inventory(main, src, profile), encoding="utf-8")

    if verbose:
        print("Phase 1 real merge...")
    # 旧 `clinical_recover_sources.py` 在模块级 `warnings.filterwarnings("ignore",
    # category=FutureWarning)`：pandas≥2.1 对 `base.at[i,"name_cn"] = <str>`
    # （float64 列赋字符串）会发 dtype 变更告警。为复现"SSD 链不出该告警"的既有表现，
    # 这里只包住 phase1，不影响 lxj_AN 链与库调用方的告警设置。
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=FutureWarning)
        real_df, conflict_df = phase1_real_merge(main, src, profile)
    meta = real_df.attrs.get("meta", {})
    export_real = real_df.drop(columns=["_name_key"], errors="ignore")
    if len(conflict_df):
        conflict_df = conflict_df.drop_duplicates(
            subset=["name_en", "field", "value_a", "value_b", "source_b"], keep="first")
    if verbose:
        print(f"  real n={len(real_df)} conflicts={len(conflict_df)} "
              f"orphans={meta.get('n_orphans')}")
        print(f"  cohort_final: {real_df['cohort_final'].value_counts().to_dict()}")
        print(f"  fill_counts: {meta.get('fill_counts')}")

    meta_df = pd.DataFrame([
        {"key": "n_rows", "value": len(real_df)},
        {"key": "n_conflicts", "value": len(conflict_df)},
        {"key": "n_orphans", "value": meta.get("n_orphans", 0)},
        {"key": "seed", "value": profile.seed},
        {"key": "fill_counts", "value": str(meta.get("fill_counts"))},
        {"key": "cohort_final_counts",
         "value": str(real_df["cohort_final"].value_counts().to_dict())},
    ])
    real_path = out / profile.out_real
    with pd.ExcelWriter(real_path, engine="openpyxl") as xw:
        export_real.to_excel(xw, sheet_name=profile.main_sheet_name, index=False)
        if len(conflict_df):
            conflict_df.to_excel(xw, sheet_name="conflict_log", index=False)
        else:
            pd.DataFrame({"note": ["no conflicts logged"]}).to_excel(
                xw, sheet_name="conflict_log", index=False)
        meta_df.to_excel(xw, sheet_name="meta", index=False)

    if verbose:
        print(f"DONE (merge) → {out}")
    return {
        "out": out, "real_path": real_path, "real": real_df,
        "conflicts": conflict_df, "meta": meta, "source_long": src, "report": out / profile.out_inventory,
    }


# ═════════════════════════════════════════════════════════════════
# 策略 3：modality_outer（DTI ⨝ IVIM）
# ═════════════════════════════════════════════════════════════════
def _values_disagree(a: Any, b: Any) -> bool:
    """模态外连接专用：两边都非缺失且数值/文本不一致（严格 atol=1e-9）。"""
    if pd.isna(a) and pd.isna(b):
        return False
    if pd.isna(a) or pd.isna(b):
        return False
    try:
        fa, fb = float(a), float(b)
        if np.isfinite(fa) and np.isfinite(fb):
            return not np.isclose(fa, fb, rtol=0.0, atol=1e-9, equal_nan=True)
    except (TypeError, ValueError):
        pass
    return str(a).strip() != str(b).strip()


def _coalesce_prefer_left(left: Any, right: Any) -> Any:
    return left if pd.notna(left) else right


def outer_merge_modalities(
    left: pd.DataFrame,
    right: pd.DataFrame,
    left_tag: str,
    right_tag: str,
    img_sep: str = "__",
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """两个模态表按 ``Name`` 外连接（原 ``outer_merge_modalities``）。

    共享临床列：左表非空优先；影像列（列名含 ``img_sep``）两边不许重名。
    """
    L = left.copy()
    R = right.copy()
    L["Name"] = L["Name"].astype(str).str.strip()
    R["Name"] = R["Name"].astype(str).str.strip()

    def _is_img(c: Any) -> bool:
        return img_sep in str(c)

    l_img = [c for c in L.columns if _is_img(c)]
    r_img = [c for c in R.columns if _is_img(c)]
    img_overlap = set(l_img) & set(r_img)
    if img_overlap:
        raise ValueError(f"imaging column name collision: {sorted(img_overlap)[:5]} ...")

    l_meta = [c for c in L.columns if c not in l_img]
    r_meta = [c for c in R.columns if c not in r_img]
    shared_meta = [c for c in l_meta if c in r_meta and c != "Name"]
    l_only_meta = [c for c in l_meta if c not in r_meta and c != "Name"]
    r_only_meta = [c for c in r_meta if c not in l_meta and c != "Name"]
    overlap_cols = [c for c in L.columns if c in R.columns and c != "Name"]

    m = L.merge(R, on="Name", how="outer", suffixes=("_L", "_R"), indicator=True)

    qc_rows: List[Dict[str, Any]] = []
    parts: List[pd.DataFrame] = []

    parts.append(pd.DataFrame({
        "Name": m["Name"],
        f"has_{left_tag}": m["_merge"].isin(["both", "left_only"]).astype(int),
        f"has_{right_tag}": m["_merge"].isin(["both", "right_only"]).astype(int),
    }))

    shared_data: Dict[str, List[Any]] = {}
    for col in shared_meta:
        if col in overlap_cols:
            s_l, s_r = m[f"{col}_L"], m[f"{col}_R"]
        else:
            s_l = m[col]
            s_r = m[col]
        chosen: List[Any] = []
        for i in range(len(m)):
            a, b = s_l.iloc[i], s_r.iloc[i]
            if _values_disagree(a, b):
                qc_rows.append({
                    "Name": m["Name"].iloc[i], "field": col,
                    f"value_{left_tag}": a, f"value_{right_tag}": b,
                    "chosen": a, "rule": f"prefer_{left_tag}_on_conflict",
                })
            chosen.append(_coalesce_prefer_left(a, b))
        shared_data[col] = chosen
    if shared_data:
        parts.append(pd.DataFrame(shared_data))

    for cols, tag in ((l_only_meta, "L"), (r_only_meta, "R")):
        block = {}
        for col in cols:
            src = f"{col}_{tag}" if f"{col}_{tag}" in m.columns else col
            block[col] = m[src]
        if block:
            parts.append(pd.DataFrame(block))

    for cols, tag in ((l_img, "L"), (r_img, "R")):
        block = {}
        for col in cols:
            src = f"{col}_{tag}" if f"{col}_{tag}" in m.columns else col
            block[col] = m[src]
        if block:
            parts.append(pd.DataFrame(block))

    out = pd.concat(parts, axis=1)
    out = out.sort_values("Name", kind="mergesort").reset_index(drop=True)

    qc = pd.DataFrame(qc_rows)
    if qc.empty:
        qc = pd.DataFrame(columns=[
            "Name", "field", f"value_{left_tag}", f"value_{right_tag}", "chosen", "rule"])
    return out, qc


# ═════════════════════════════════════════════════════════════════
# CLI
# ═════════════════════════════════════════════════════════════════
DEFAULT_STRATEGY = {"lxj_AN": "keyed_union", "ssd": "source_records"}


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        prog="tidy_merge",
        description="统一合并入口：keyed_union(lxj_AN) / source_records(ssd) / modality_outer",
    )
    ap.add_argument("--profile", help="lxj_AN | ssd")
    ap.add_argument("--strategy", choices=["auto", "keyed_union", "source_records", "modality_outer"],
                    default="auto")
    ap.add_argument("--out", help="keyed_union：输出工作簿路径")
    ap.add_argument("--out-dir", help="source_records：输出目录")
    ap.add_argument("--left", help="modality_outer：左表 xlsx")
    ap.add_argument("--right", help="modality_outer：右表 xlsx")
    ap.add_argument("--left-tag", default="dti")
    ap.add_argument("--right-tag", default="ivim")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args(argv)

    if not args.profile and args.strategy != "modality_outer":
        ap.error("--profile is required (or use --strategy modality_outer)")
    profile = get_profile(args.profile) if args.profile else None

    strategy = args.strategy
    if strategy == "auto":
        strategy = DEFAULT_STRATEGY.get(profile.name)
        if strategy is None:
            ap.error(f"no default strategy for profile {profile.name}; pass --strategy")

    if strategy == "keyed_union":
        merge_keyed_union(profile, Path(args.out) if args.out else None, verbose=not args.quiet)
        return 0

    if strategy == "source_records":
        merge_source_records(profile, Path(args.out_dir) if args.out_dir else None,
                             verbose=not args.quiet)
        return 0

    # modality_outer
    if not (args.left and args.right):
        ap.error("--left and --right are required for modality_outer")
    left = read_frame(args.left, sheet=0)
    right = read_frame(args.right, sheet=0)
    merged, qc = outer_merge_modalities(left, right, args.left_tag, args.right_tag)
    if args.out:
        write_frames(Path(args.out), collections.OrderedDict(
            [("merged", merged), ("merge_qc", qc)]))
        print(f"written → {args.out} (rows={len(merged)}, qc={len(qc)})")
    else:
        print(merged.head().to_string())
        print(f"rows={len(merged)} qc={len(qc)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
