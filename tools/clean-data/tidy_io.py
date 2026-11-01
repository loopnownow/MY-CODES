# -*- coding: utf-8 -*-
"""tidy_io — Excel 读写与规范化公共层。

合并来源：
  * `lxj_AN/build_merge.py`  的 ``cellval`` / ``norm`` / ``load`` / ``style_header``
  * `lxj_AN/build_complete.py` 的 ``nz`` / ``s``
  * `lxj_AN/verify_full.py`   的同名读取函数
  * `ssd/clinical_recover_sources.py` 的姓名/性别/数值规范化与 `_set`

两类读取路径刻意保留差异，不要互相替换：

``XlsBook``  (``.xls``)
    用 xlrd 逐格读取，保留日期/文本类型语义——lxj_AN 三个源文件是 xls，
    其日期列必须走 ``XL_CELL_DATE`` 分支才能得到 ``YYYY-MM-DD``。
``read_frame`` (``.xlsx`` / ``.xls``)
    用 pandas 读成 DataFrame——SSD 侧源文件是 xlsx，列名映射按 DataFrame 走。
"""
from __future__ import annotations

import collections
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple, Union

import numpy as np
import pandas as pd
import xlrd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ═════════════════════════════════════════════════════════════════
# 单元格值规范化
# ═════════════════════════════════════════════════════════════════
def nz(v: Any) -> bool:
    """非空判定（None / 空白串 视为空）。"""
    return v is not None and str(v).strip() != ""


def s(v: Any) -> str:
    """单元格转字符串；None → ''。

    不可用 ``str(None)``——那会变成字面量 ``'None'`` 而绕过所有空值判断。
    """
    return "" if v is None else str(v).strip()


def norm(v: Any) -> str:
    """数值化字符串：整数去小数点，其余 strip。"""
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def to_num(v: Any) -> float:
    """尽力转 float，失败返回 nan；会剥掉尾部中文单位/备注。"""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return np.nan
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    text = str(v).strip()
    if not text or text.lower() in {"nan", "none", "null", "na", "-", "/"}:
        return np.nan
    text = re.sub(r"[^\d.\-]+$", "", text)
    try:
        return float(text)
    except ValueError:
        return np.nan


def norm_name(x: Any) -> Optional[str]:
    """英文姓名规范化：大写、去分隔符、去非字母数字。"""
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return None
    text = str(x).strip().upper()
    if not text or text in {"NAN", "NONE", "NULL", "NA"}:
        return None
    text = text.replace("_", " ").replace("-", " ")
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^A-Z0-9 ]", "", text)
    return text.strip() or None


def name_key_compact(x: Any) -> Optional[str]:
    """姓名紧凑键（去所有空白），用于跨源对齐。"""
    n = norm_name(x)
    return None if not n else re.sub(r"\s+", "", n)


def normalize_sex(v: Any) -> Optional[str]:
    """性别规范化为 ``'male'`` / ``'female'``；无法判定返回 None。"""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    text = str(v).strip().lower()
    if text in {"", "nan", "none", "null"}:
        return None
    try:
        f = float(text)
        if f == 1:
            return "male"
        if f == 2:
            return "female"
    except ValueError:
        pass
    if text in {"m", "male", "man", "男", "1.0"}:
        return "male"
    if text in {"f", "female", "woman", "女", "2.0"}:
        return "female"
    if "男" in str(v):
        return "male"
    if "女" in str(v):
        return "female"
    return None


def values_conflict(a: Any, b: Any, tol: float = 1e-6) -> bool:
    """两边都非空且实质不同 → True（数值按相对容差比较）。"""
    if a is None or b is None:
        return False
    if isinstance(a, float) and np.isnan(a):
        return False
    if isinstance(b, float) and np.isnan(b):
        return False
    na, nb = to_num(a), to_num(b)
    if not np.isnan(na) and not np.isnan(nb):
        return abs(na - nb) > max(tol, 0.01 * max(abs(na), abs(nb), 1.0))
    return str(a).strip().lower() != str(b).strip().lower()


def is_blank(v: Any) -> bool:
    """DataFrame 语义下的空（None / nan / 空白串）。"""
    if v is None:
        return True
    if isinstance(v, float) and np.isnan(v):
        return True
    return isinstance(v, str) and not v.strip()


def set_if(rec: Dict[str, Any], key: str, val: Any, numeric: bool = False) -> None:
    """把非空 ``val`` 写入 ``rec[key]``（合并自 `_set`）。"""
    if is_blank(val):
        return
    if numeric:
        num = to_num(val)
        if np.isnan(num):
            return
        rec[key] = num
    else:
        rec[key] = val


# ═════════════════════════════════════════════════════════════════
# .xls 逐格读取（xlrd）
# ═════════════════════════════════════════════════════════════════
class XlsBook:
    """xlrd 工作簿薄封装，提供与原脚本一致的 ``cellval`` 语义。"""

    def __init__(self, path: Union[str, Path]):
        self.path = Path(path)
        self.book = xlrd.open_workbook(str(self.path))
        self.datemode = self.book.datemode
        self.sheet_names: List[str] = list(self.book.sheet_names())

    def sheet(self, name_or_index: Union[str, int]) -> Any:
        if isinstance(name_or_index, int):
            return self.book.sheet_by_index(name_or_index)
        return self.book.sheet_by_name(name_or_index)

    def cellval(self, sh: Any, r: int, c: int) -> Any:
        """读一格，返回干净 Python 值（与原 ``cellval`` 等价）。"""
        if c >= sh.ncols:
            return ""
        t = sh.cell_type(r, c)
        v = sh.cell_value(r, c)
        if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return ""
        if t == xlrd.XL_CELL_TEXT:
            return v.strip()
        if t == xlrd.XL_CELL_NUMBER:
            return int(v) if float(v).is_integer() else v
        if t == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate_as_datetime(v, self.datemode).strftime("%Y-%m-%d")
            except Exception:
                return v
        if t == xlrd.XL_CELL_BOOLEAN:
            return "TRUE" if v else "FALSE"
        return str(v).strip()

    def header(self, sh: Any, hdr_row: int = 0) -> List[str]:
        return [norm(self.cellval(sh, hdr_row, c)) for c in range(sh.ncols)]

    def rows(self, sh: Any, hdr_row: int = 0) -> List[List[Any]]:
        """表头之后的所有行（原始值，已过 ``cellval``）。"""
        out = []
        for r in range(hdr_row + 1, sh.nrows):
            out.append([self.cellval(sh, r, c) for c in range(sh.ncols)])
        return out


def open_book(path: Union[str, Path]) -> XlsBook:
    """打开 .xls。.xlsx 请走 :func:`read_frame`。"""
    p = Path(path)
    if p.suffix.lower() != ".xls":
        raise ValueError(f"open_book expects .xls, got {p.name}; use read_frame for .xlsx")
    return XlsBook(p)


def load_keyed(
    bk: XlsBook,
    sh: Any,
    key_col: int,
    hdr_row: int = 0,
    skip_keys: Iterable[str] = (),
    sheet_name: Optional[str] = None,
) -> Tuple[List[str], Dict[str, List[Any]], List[Tuple[str, str, int, str]], List[str]]:
    """按主键装载一张表，同键多行取「非空字段最多」的一行。

    与 `lxj_AN/build_merge.py` 的 ``load()`` 逐行等价。

    Returns
    -------
    (header, rows_by_key, dup_log, key_order)
        ``dup_log`` 每项为 ``(sheet, key, 出现次数, '完全相同' | '内容有冲突')``。
    """
    skip_keys = set(skip_keys)
    name = sheet_name if sheet_name is not None else getattr(sh, "name", "?")
    hdr = [norm(bk.cellval(sh, hdr_row, c)) for c in range(sh.ncols)]
    buckets: Dict[str, List[List[Any]]] = collections.defaultdict(list)
    order: List[str] = []
    for r in range(hdr_row + 1, sh.nrows):
        k = norm(bk.cellval(sh, r, key_col))
        if not k or k in skip_keys:
            continue
        row = [bk.cellval(sh, r, c) for c in range(sh.ncols)]
        if not any(x != "" for x in row):
            continue
        if k not in buckets:
            order.append(k)
        buckets[k].append(row)

    final: Dict[str, List[Any]] = {}
    dups: List[Tuple[str, str, int, str]] = []
    for k, rows in buckets.items():
        if len(rows) > 1:
            best = max(rows, key=lambda rw: sum(1 for x in rw if x != ""))
            identical = len({tuple(norm(x) for x in rw) for rw in rows}) == 1
            dups.append((name, k, len(rows), "完全相同" if identical else "内容有冲突"))
            final[k] = best
        else:
            final[k] = rows[0]
    return hdr, final, dups, order


# ═════════════════════════════════════════════════════════════════
# .xlsx / DataFrame 读取
# ═════════════════════════════════════════════════════════════════
def sheet_names(path: Union[str, Path]) -> List[str]:
    """工作簿表名（xls / xlsx 均可）。"""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    if p.suffix.lower() == ".xls":
        return XlsBook(p).sheet_names
    with pd.ExcelFile(p) as xl:
        return list(xl.sheet_names)


def read_frame(
    path: Union[str, Path],
    sheet: Union[str, int, None] = 0,
    required: bool = True,
    **kwargs: Any,
) -> Optional[pd.DataFrame]:
    """读一张表为 DataFrame；``required=False`` 且表不存在时返回 None。"""
    p = Path(path)
    if not p.exists():
        if required:
            raise FileNotFoundError(p)
        return None
    if isinstance(sheet, str):
        names = sheet_names(p)
        if sheet not in names:
            if required:
                raise KeyError(f"{p.name}: sheet {sheet!r} not found (have {names})")
            return None
    return pd.read_excel(p, sheet_name=sheet, **kwargs)


def write_frames(
    path: Union[str, Path],
    frames: "collections.OrderedDict[str, pd.DataFrame]",
    index: bool = False,
) -> Path:
    """把多个 DataFrame 写成多表工作簿。"""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(p, engine="openpyxl") as xw:
        for name, df in frames.items():
            df.to_excel(xw, sheet_name=name[:31], index=index)
    return p


# ═════════════════════════════════════════════════════════════════
# openpyxl 写出样式（lxj_AN 侧风格）
# ═════════════════════════════════════════════════════════════════
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
KEY_FILL = PatternFill("solid", fgColor="DDEBF7")


def style_header(ws: Any, ncol: int, widths: Optional[Sequence[float]] = None) -> None:
    """表头底色/字体/冻结/筛选，可选列宽。"""
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{ws.max_row}"
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def mark_key_column(ws: Any, col: int = 1, width: float = 15) -> None:
    """首列（编号）淡蓝加粗 + 列宽。"""
    ws.column_dimensions[get_column_letter(col)].width = width
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=col)
        cell.fill = KEY_FILL
        cell.font = Font(bold=True)
