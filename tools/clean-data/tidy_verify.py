#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""tidy_verify — 统一「校验」入口，两种策略对应原来的两套校验代码。

策略
----
``cell``  逐格回查：主表每个单元格必须等于某个源行里的对应值；另做结构断言
          （编号唯一、分表行序与主表一致、并集数独立重算）。
          原 `lxj_AN/verify_full.py`（profile ``lxj_AN``）。
``paper`` 三层数据（真实合并 / 论文感知填充 / 纯合成）与论文表 1 的 mean±SD 对照，
          产出 `validation_vs_paper.md` 与 `README_clinical_recovery.md`。
          原 `ssd/clinical_recover_report.py`（profile ``ssd``）。

用法::

    python tidy_verify.py --profile lxj_AN
    python tidy_verify.py --profile ssd
    python tidy_verify.py --profile lxj_AN --fail-on-diff   # 有差异时退出码 1
"""
from __future__ import annotations

import argparse
import collections
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from tidy_io import XlsBook, norm, open_book  # noqa: E402
from tidy_profiles import LxjAnProfile, SsdProfile, get_profile  # noqa: E402

__all__ = [
    "verify_cells",
    "verify_paper",
    "summarize_group",
    "fmt_paper",
    "write_validation",
    "write_readme",
    "main",
]


# ═════════════════════════════════════════════════════════════════
# 策略 1：cell（逐格回查）
# ═════════════════════════════════════════════════════════════════
def _sheet_multi(bk: XlsBook, sh: Any, key_col: int, hdr_row: int = 0):
    """按主键收集**全部**原始行（不去重），供"等于任一源值"判定。"""
    hd = [norm(bk.cellval(sh, hdr_row, c)) for c in range(sh.ncols)]
    m: Dict[str, List[List[Any]]] = collections.defaultdict(list)
    for r in range(hdr_row + 1, sh.nrows):
        k = norm(bk.cellval(sh, r, key_col))
        if k:
            m[k].append([norm(bk.cellval(sh, r, c)) for c in range(sh.ncols)])
    return hd, m


def verify_cells(
    profile: LxjAnProfile,
    workbook: Optional[Path] = None,
    verbose: bool = True,
) -> Dict[str, Any]:
    """独立逐格校验：主表单元格 vs 三个源 .xls。等价于 `lxj_AN/verify_full.py`。"""
    p = profile
    out_path = Path(workbook) if workbook else p.root / p.out_merge

    bk_a = open_book(p.path(p.file_a))
    bk_b = open_book(p.path(p.file_b))
    bk_c = open_book(p.path(p.file_c))

    wb = load_workbook(out_path, read_only=True)
    ws = wb[p.master_sheet]
    allrows = list(ws.iter_rows(values_only=True))
    hdr = [norm(x) for x in allrows[0]]
    idx = {h: i for i, h in enumerate(hdr)}
    master = {norm(r[0]): r for r in allrows[1:]}
    if verbose:
        print(f"{p.master_sheet} {len(allrows) - 1} 行 x {len(hdr)} 列")

    fails: List[str] = []
    checks = 0

    def expect_at_least_one(key: str, outcol: str, values: Sequence[Any], src: str) -> None:
        """主表单元格必须等于源中（同键多行）**任一**值——去重规则保留的是最完整一行。"""
        nonlocal checks
        if key not in master or outcol not in idx:
            return
        got = norm(master[key][idx[outcol]])
        checks += 1
        allowed = [norm(v) for v in values]
        if got not in allowed:
            fails.append(f'{src} {key} [{outcol}] master="{got}" '
                         f"not in source values {sorted(set(allowed))}")

    # ── B ──
    shB = bk_b.sheet(0)
    hdB, mB = _sheet_multi(bk_b, shB, 0)
    for k, rows in mB.items():
        if k == "1":
            continue
        for c in range(4, shB.ncols):
            expect_at_least_one(k, hdB[c], [rw[c] for rw in rows], "B")
        for c, out in p.verify_b_cols:
            if norm(rows[0][c]):
                expect_at_least_one(k, out, [rw[c] for rw in rows], "B")

    # ── C ──
    shC = bk_c.sheet(0)
    hdC, mC = _sheet_multi(bk_c, shC, 1)
    for k, rows in mC.items():
        if k == "1":
            continue
        for c, out in p.verify_c_cols:
            expect_at_least_one(k, out, [rw[c] for rw in rows], "C")
        for c in range(7, 17):
            expect_at_least_one(k, "C_" + hdC[c], [rw[c] for rw in rows], "C")
        if norm(rows[0][2]):
            expect_at_least_one(k, "姓名", [rw[2] for rw in rows], "C")

    # ── A（7 张表） ──
    pfx = dict(p.verify_prefix)
    for sname, kc, cmap in p.verify_sheets:
        sh = bk_a.sheet(sname)
        hd, m = _sheet_multi(bk_a, sh, kc)
        for k, rows in m.items():
            for c, out in cmap.items():
                name = out if out != "@" else pfx[sname] + hd[c]
                expect_at_least_one(k, name, [rw[c] for rw in rows], "A/" + sname)
            # 文化程度 / 是否为右利手：各表列位不同，按表头名定位。
            # 注意 profile.field_names 是 (主表列名, 源表头名)；旧 build_merge 用同一元组时
            # 顺序相反，这里必须按 (out, want) 解包，否则会漏掉「是否为右利手」整列检查。
            for out, want in p.field_names:
                if want in hd:
                    c = hd.index(want)
                    if norm(rows[0][c]):
                        expect_at_least_one(k, out, [rw[c] for rw in rows],
                                            "A/" + sname + "/" + want)

    if verbose:
        print(f"\n逐格核对完成: 共 {checks} 项检查, 失败 {len(fails)} 项")
        for f in fails[:30]:
            print("   FAIL", f)

    # ── 结构断言 ──
    order = [norm(r[0]) for r in allrows[1:]]
    structural_ok = True
    if verbose:
        print(f"\n编号唯一={len(set(order)) == len(order)}  行数={len(order)}")
    if len(set(order)) != len(order):
        structural_ok = False
    for name in wb.sheetnames:
        if name in (p.master_sheet, p.log_sheet):
            continue
        w = wb[name]
        rr = list(w.iter_rows(values_only=True))
        keys = [norm(x[0]) for x in rr[1:]]
        if not (len(keys) == len(order) and all(a == b for a, b in zip(keys, order))):
            structural_ok = False
            if verbose:
                print(f"   !! 分表 {name} 行数/行序与主表不一致")
    if verbose:
        print(f"所有分表行数={len(order)} 且行序与总表完全一致: {structural_ok}")

    # ── 并集数独立重算 ──
    all_a: set = set()
    for sname, kc, _ in p.verify_sheets:
        sh = bk_a.sheet(sname)
        for r in range(1, sh.nrows):
            v = norm(bk_a.cellval(sh, r, kc))
            if v:
                all_a.add(v)
    core = {k for k in mB if k != "1"} | {k for k in mC if k != "1"}
    if verbose:
        print(f"独立重算并集={len(all_a | core)}  "
              f"(A={len(all_a)} core={len(core)} 交={len(all_a & core)})")

    return {
        "checks": checks, "fails": fails, "structural_ok": structural_ok,
        "n_rows": len(order), "union_recount": len(all_a | core), "path": out_path,
    }


# ═════════════════════════════════════════════════════════════════
# 策略 2：paper（与论文表 1 对照）
# ═════════════════════════════════════════════════════════════════
def summarize_group(df: pd.DataFrame, group_col: str, gname: str,
                    vars_: Sequence[str]) -> Dict[str, str]:
    sub = df[df[group_col] == gname]
    out: Dict[str, str] = {"n": str(len(sub))}
    if "sex" in sub.columns:
        out["sex_M/F"] = f"{(sub['sex'] == 'male').sum()}/{(sub['sex'] == 'female').sum()}"
    for v in vars_:
        if v not in sub.columns:
            out[v] = "—"
            continue
        s = pd.to_numeric(sub[v], errors="coerce").dropna()
        out[v] = "n=0" if len(s) == 0 else f"n={len(s)} {s.mean():.2f}±{s.std():.2f}"
    return out


def fmt_paper(info: dict, vars_: Sequence[str]) -> Dict[str, str]:
    out: Dict[str, str] = {"n": str(info["n"]), "sex_M/F": f"{info['male']}/{info['female']}"}
    for v in vars_:
        if v in info and isinstance(info[v], tuple):
            m, s = info[v]
            out[v] = f"{m:.2f}±{s:.2f}"
        else:
            out[v] = "—"
    return out


def write_validation(real_df: pd.DataFrame, filled_df: pd.DataFrame,
                     syn1: pd.DataFrame, syn2: pd.DataFrame,
                     profile: SsdProfile) -> str:
    """三层数据 vs 论文表 1 校验报告（原 ``write_validation``）。"""
    lines: List[str] = []
    lines.append("# 校验报告：真实 / 填充 / 合成 vs 论文表1\n\n")
    vars_p2 = ["age", "education_years", "mmse", "phq15", "sss", "gad7", "hamd"]

    def table_block(title: str, paper_side: dict, real_s: dict,
                    filled_s: dict, syn_s: Optional[dict] = None) -> None:
        lines.append(f"### {title}\n\n")
        lines.append("| 指标 | 论文 | 真实合并 | 论文感知填充 | 纯合成 |\n")
        lines.append("|------|------|----------|--------------|--------|\n")
        for k in ["n", "sex_M/F"] + vars_p2:
            lines.append(
                f"| {k} | {paper_side.get(k, '—')} | {real_s.get(k, '—')} | "
                f"{filled_s.get(k, '—')} | {(syn_s.get(k, '—') if syn_s else '—')} |\n"
            )
        lines.append("\n")

    lines.append("## 论文2 对照（推荐主校验）\n\n")
    for g in ["SSD", "HC"]:
        table_block(
            f"论文2 {g}",
            fmt_paper(profile.paper2[g], vars_p2),
            summarize_group(real_df, "cohort_final", g, vars_p2),
            summarize_group(filled_df, "cohort_final", g, vars_p2),
            summarize_group(syn2, "cohort_final", g, vars_p2),
        )

    lines.append("## 论文1 对照\n\n")
    vars_p1 = ["age", "education_years", "phq15", "sss"]
    for g in ["SSD", "HC"]:
        table_block(
            f"论文1 {g}",
            fmt_paper(profile.paper1[g], vars_p1),
            summarize_group(real_df, "cohort_final", g, vars_p1),
            summarize_group(filled_df, "cohort_final", g, vars_p1),
            summarize_group(syn1, "cohort_final", g, vars_p1),
        )

    lines.append("## 真实合并后关键字段完整度（cohort_final）\n\n")
    lines.append("| 字段 | SSD | HC | unknown |\n|------|-----|----|----------|\n")
    for c in ["age", "sex", "education_years", "mmse", "phq15", "sss", "gad7", "hamd"]:
        cells = []
        for g in ["SSD", "HC", "unknown"]:
            sub = real_df[real_df["cohort_final"] == g]
            nn = sub[c].notna().sum() if c in sub.columns else 0
            cells.append(f"{nn}/{len(sub)}")
        lines.append(f"| {c} | {cells[0]} | {cells[1]} | {cells[2]} |\n")

    lines.append("\n## 使用建议\n\n")
    lines.append("- **发表/真实分析**: 使用 `clinical_real_merged.xlsx`（无任何合成），或填充表中 `any_synthetic_core==0` 子集（允许仅 HAMD 为合成）。\n")
    lines.append("- **方法演示/pipeline 试跑**: 可用填充全表或纯合成表；须声明 SYNTHETIC。\n")
    lines.append("- **HAMD**: 真实源无数据；填充表中全员 `is_synthetic_hamd=1`，故 `any_synthetic` 几乎恒为 1；请用 `any_synthetic_core` 判断核心变量是否合成。\n")
    if "any_synthetic_core" in filled_df.columns:
        lines.append(
            f"\n- 填充表 `any_synthetic_core==1` 行数: "
            f"{int(filled_df['any_synthetic_core'].sum())}/{len(filled_df)}\n")
        lines.append(
            f"- 填充表 `real_core` 行数: "
            f"{int((filled_df['any_synthetic_core'] == 0).sum())}/{len(filled_df)}\n")
    return "".join(lines)


README_TEMPLATE = """# 临床资料恢复包说明 (README)

## 文件

| 文件 | 含义 |
|------|------|
| `{inventory}` | 盘点：来源、缺失、与论文对照 |
| `{real}` | **仅真实值**合并 + cohort 回推 + conflict_log |
| `{filled}` | 真实 + 按论文 mean±SD 补缺（带 `is_synthetic_*`；sheet `real_core_subset`） |
| `{fill_summary}` | 各变量合成填充计数 |
| `{syn1}` | 论文1 表1 整队列合成（101+36） |
| `{syn2}` | 论文2 表1-1 整队列合成（56+46，含 HAMD） |
| `{validation}` | 三层数据 vs 论文校验 |
| `source_long.parquet` / `source_long.csv` | 多源长表（调试） |

## 工作表 (real / filled)

- `clinical`: 主表
- `conflict_log`: 多源冲突（保留主库值）
- `meta`: 合并统计

## 关键列

- `cohort_final` / `cohort_rule`: 最终分组与规则
- `sources_merged` / `n_sources`: 来源追溯
- `is_synthetic_<var>`: 该字段是否合成
- `any_synthetic`: 含 HAMD 的任意合成（因 HAMD 全合成，几乎恒为 1）
- `any_synthetic_core` / `real_core_complete`: **不含 HAMD** 的核心变量合成标记（推荐用这个筛选）
- `hamd`: 仅合成路径有值

## 方法学声明模板

> 原始逐例临床表存在缺失。我们首先合并全部现存登记表与名单，保留真实观测值并记录冲突。对分析所需连续变量的缺失值，按已发表基线表（mean±SD）与固定随机种子生成合成补缺，并以 `is_synthetic_*` 标记。发表性分析仅使用 `clinical_real_merged.xlsx` 或 `any_synthetic_core==0` 子集。HAMD 无真实源，不得当作实测。纯合成队列仅用于方法演示，不代表真实患者。

## 不要做的事

- 不要把合成 HAMD/GAD 当作实测写入论文结果
- 不要期望本包恢复 rs-fMRI/VBM 原始影像
"""


def write_readme(profile: Optional[SsdProfile] = None) -> str:
    p = profile
    if p is None:
        return README_TEMPLATE
    return README_TEMPLATE.format(
        inventory=p.out_inventory, real=p.out_real, filled=p.out_filled,
        fill_summary=p.out_fill_summary, syn1=p.out_syn1, syn2=p.out_syn2,
        validation=p.out_validation, source_long=p.out_source_long,
    )


def verify_paper(profile: SsdProfile, out_dir: Optional[Path] = None,
                 verbose: bool = True) -> Dict[str, Any]:
    """三层数据 vs 论文表 1 → validation_vs_paper.md + README（原 Phase 4）。"""
    out = Path(out_dir) if out_dir else profile.out
    needed = {
        "真值合并表": out / profile.out_real,
        "论文感知填充表": out / profile.out_filled,
        "纯合成表1": out / profile.out_syn1,
        "纯合成表2": out / profile.out_syn2,
    }
    missing = [f"{k}({v})" for k, v in needed.items() if not v.exists()]
    if missing:
        raise FileNotFoundError(
            "缺少输入: " + "; ".join(missing) +
            "；请先跑 `tidy_merge.py --profile ssd` 与 "
            "`tidy_complete.py --profile ssd --strategy paper_fill`（或用 --out-dir 指定）")
    real_df = pd.read_excel(needed["真值合并表"], sheet_name=profile.main_sheet_name)
    filled_df = pd.read_excel(needed["论文感知填充表"], sheet_name=profile.main_sheet_name)
    syn1 = pd.read_excel(needed["纯合成表1"])
    syn2 = pd.read_excel(needed["纯合成表2"])

    val = write_validation(real_df, filled_df, syn1, syn2, profile)
    (out / profile.out_validation).write_text(val, encoding="utf-8")
    (out / profile.out_readme).write_text(write_readme(profile), encoding="utf-8")
    if verbose:
        print(f"validation → {out / profile.out_validation}")
        print(f"readme     → {out / profile.out_readme}")
    return {"validation": out / profile.out_validation,
            "readme": out / profile.out_readme, "text": val}


# ═════════════════════════════════════════════════════════════════
# CLI
# ═════════════════════════════════════════════════════════════════
DEFAULT_STRATEGY = {"lxj_AN": "cell", "ssd": "paper"}


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        prog="tidy_verify",
        description="统一校验入口：cell(lxj_AN 逐格) / paper(ssd 论文对照)",
    )
    ap.add_argument("--profile", required=True, help="lxj_AN | ssd")
    ap.add_argument("--strategy", choices=["auto", "cell", "paper"], default="auto")
    ap.add_argument("--workbook", help="cell：被校验的主表工作簿（默认 <root>/lxj_AN.xlsx）")
    ap.add_argument("--out-dir", help="paper：恢复包目录（默认 <root>/clinical_recovery）")
    ap.add_argument("--fail-on-diff", action="store_true",
                    help="有失败项/结构异常时返回退出码 1（默认沿用原脚本：始终 0）")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args(argv)

    profile = get_profile(args.profile)
    strategy = args.strategy
    if strategy == "auto":
        strategy = DEFAULT_STRATEGY.get(profile.name, "cell")
    verbose = not args.quiet

    if strategy == "cell":
        res = verify_cells(profile, Path(args.workbook) if args.workbook else None, verbose=verbose)
        ok = not res["fails"] and res["structural_ok"]
        if args.fail_on_diff and not ok:
            return 1
        return 0

    verify_paper(profile, Path(args.out_dir) if args.out_dir else None, verbose=verbose)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
