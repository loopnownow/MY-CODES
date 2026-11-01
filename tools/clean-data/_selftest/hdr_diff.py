# -*- coding: utf-8 -*-
"""Compare sheet HEADERS (and shapes) of two workbooks — for matching the RAD console layout."""
import sys
from openpyxl import load_workbook


def headers(path):
    wb = load_workbook(path, read_only=True)
    out = {}
    for n in wb.sheetnames:
        ws = wb[n]
        hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
        out[n] = (ws.max_row - 1, hdr)
    wb.close()
    return out


def main(a, b):
    ha, hb = headers(a), headers(b)
    rc = 0
    for n in ha:
        if n not in hb:
            print(f"[{n}] missing in B")
            rc = 1
            continue
        ra, xa = ha[n]
        rb, xb = hb[n]
        same = xa == xb
        print(f"[{n}] rows {ra} vs {rb}, cols {len(xa)} vs {len(xb)}, headers "
              f"{'SAME ✓' if same else 'DIFFER'}")
        if not same:
            rc = 1
            for i in range(max(len(xa), len(xb))):
                va = xa[i] if i < len(xa) else "<none>"
                vb = xb[i] if i < len(xb) else "<none>"
                if va != vb:
                    print(f"    col {i + 1}: A={va!r}  B={vb!r}")
        if ra != rb:
            rc = 1
    return rc


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1], sys.argv[2]))
