# -*- coding: utf-8 -*-
"""Split the DTI / IVIM sheets of a container workbook into standalone files.

Used to exercise the `tidy_merge.py --strategy modality_outer` CLI path (which reads
single-sheet workbooks), then compare its result against the verified
`tidy_split.py` output.

usage: python make_outer_inputs.py <SSD_imputed.xlsx> <outdir>
"""
import sys
from pathlib import Path

import pandas as pd

src, outdir = Path(sys.argv[1]), Path(sys.argv[2])
outdir.mkdir(parents=True, exist_ok=True)
for sheet in ("DTI", "IVIM"):
    df = pd.read_excel(src, sheet_name=sheet)
    dst = outdir / f"_{sheet}.xlsx"
    df.to_excel(dst, index=False)
    print(f"{sheet}: {df.shape} -> {dst}")

# 比对参考（tidy_split 已验证与旧脚本逐格相同）
reference = Path(sys.argv[3]) if len(sys.argv) > 3 else None
if reference and reference.exists():
    from openpyxl import load_workbook

    wb = load_workbook(reference, read_only=True)
    for n in wb.sheetnames:
        print("reference sheet", n, wb[n].max_row, "x", wb[n].max_column)
    wb.close()
