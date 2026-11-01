# -*- coding: utf-8 -*-
"""Compare two xlsx workbooks cell-by-cell (all sheets, all cells)."""
import sys
from openpyxl import load_workbook


def cells(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(tuple("" if v is None else str(v) for v in row))
        out[name] = rows
    wb.close()
    return out


def main(a, b, max_report=20):
    ca, cb = cells(a), cells(b)
    problems = []
    if list(ca) != list(cb):
        problems.append(f"sheet names differ:\n  A={list(ca)}\n  B={list(cb)}")
    for name in ca:
        if name not in cb:
            continue
        ra, rb = ca[name], cb[name]
        if len(ra) != len(rb):
            problems.append(f"[{name}] row count {len(ra)} vs {len(rb)}")
        n = 0
        for i, (x, y) in enumerate(zip(ra, rb)):
            if x != y:
                n += 1
                if len(problems) < max_report:
                    diffs = [(j, xx, yy) for j, (xx, yy) in enumerate(zip(x, y)) if xx != yy]
                    problems.append(f"[{name}] row {i+1}: {len(diffs)} cell diff(s) {diffs[:5]}")
        if n:
            problems.append(f"[{name}] total differing rows: {n}")
        else:
            print(f"[{name}] OK ({len(ra)} rows x {len(ra[0]) if ra else 0} cols)")
    if problems:
        print("\nDIFFERENCES:")
        for p in problems:
            print("  " + p)
        return 1
    print("\nIDENTICAL ✓")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1], sys.argv[2]))
