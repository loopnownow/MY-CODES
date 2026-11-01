# -*- coding: utf-8 -*-
"""把随访列回填进 yt_RDL.xlsx 的 cln 表，并补入新患者行。

做三件事
--------
1. **回填**：候选表里能对上 `cln.Patient_ID`（姓名拼音首字母）且 **性别 + 年龄（±3 岁）双重吻合**
   的患者，把 10 个随访列写进 `cln`。
2. **补行**：候选表里 `cln` 尚无的患者，按 `cln` 的列结构追加到末尾；能从住院资料得到的
   `Gender` / `Age` 一并填上，其余（Dedifferentiated / Recurrence / Reader_* / KI67 / CDK4 /
   MDM2 / Surrounding involving / Tumor density / Satellite nodules）留空——这些是本课题
   自己的判读结果，不能由外部资料推断。
3. **一人多卡号**：候选表按「病历卡号+身份证号」计数，同一身份证号若出现多行（朱占勇），
   `cln` 是患者级表，只补一行，取**基线日期最早**的那行。

新增列（顺序与用户给定一致）::

    无复发生存时间(月) 复发状态 复发_6月 复发_6月_含转移 复发_1年 复发_1年_含转移
    复发_2年 复发_2年_含转移 复发_3年 复发_3年_含转移

安全措施：写回前把 `yt_RDL.xlsx` 备份到 `0del`；只改 `cln` 表，`rad`/`qc`/`icc` 原样保留。

运行：
    python yt_rdl_fill_yt_cln.py --out <测试副本.xlsx>   # 先写副本核对
    python yt_rdl_fill_yt_cln.py                        # 正式写回 yt_RDL.xlsx
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import tidy_yt_rdl as PATHS  # noqa: E402

FU_COLS = ["无复发生存时间(月)", "复发状态",
           "复发_6月", "复发_6月_含转移",
           "复发_1年", "复发_1年_含转移",
           "复发_2年", "复发_2年_含转移",
           "复发_3年", "复发_3年_含转移"]
AGE_TOL = 3


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan", "NaT")


def load_fu() -> pd.DataFrame:
    return pd.read_excel(PATHS.FU_XLSX, sheet_name="01_候选患者_合并")


def yt_sex(v) -> str:
    s = str(v).strip()
    return "Male" if s == "男" else ("Female" if s == "女" else "")


def confidence(row, yt_pid2age: dict) -> tuple[str, str]:
    """返回 (可信度, 说明)。可信度 ∈ 高/中/低/无。"""
    pid = str(row["拼音首字母"]).strip()
    if pid not in yt_pid2age:
        return "无", "首字母不在 cln"
    s_ok = row.get("性别一致")
    d = row.get("年龄差(岁)")
    if str(row.get("匹配可信度", "")).startswith("高"):
        return "高", "性别一致且年龄差≤3岁"
    if str(row.get("匹配可信度", "")).startswith("中"):
        return "中", f"性别一致但年龄差{'' if _blank(d) else int(d)}岁（疑似同名同首字母的他人）"
    if str(s_ok) == "否":
        return "低", "性别不符"
    return "中", "无法判定性别/年龄"


def pick_person_rows(fu: pd.DataFrame, xlsx: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """拆成 (回填行, 新增行, 一人多卡号被合并掉的行)。

    ``xlsx`` 必须是**本次要写的那个文件**——否则对副本操作时会拿原表判断谁是新患者，
    导致重复追加。
    """
    yt = pd.read_excel(xlsx, sheet_name="cln")   # 必须用全表：判断“谁是新患者”
    yt_ids = yt["Patient_ID"].astype(str).str.strip()
    yt_pid2age = dict(zip(yt_ids, pd.to_numeric(yt["Age"], errors="coerce")))

    fu = fu.copy()
    fu["_ini"] = fu["拼音首字母"].astype(str).str.strip()
    fu["_conf"], fu["_conf_note"] = zip(*[confidence(r, yt_pid2age) for _, r in fu.iterrows()])

    matched = fu[fu["_ini"].isin(set(yt_ids))].copy()
    new = fu[~fu["_ini"].isin(set(yt_ids))].copy()

    # 一人多卡号：同一身份证号（空则用病历卡号）只留基线最早的一行
    def person_key(r):
        i = "" if _blank(r["身份证号"]) else re.sub(r"[\s\u3000']", "", str(r["身份证号"]))
        return i or re.sub(r"[\s\u3000]", "", str(r["病历卡号"]))

    new["_pk"] = [person_key(r) for _, r in new.iterrows()]
    new["_base"] = pd.to_datetime(new["基线日期"], errors="coerce")
    new = new.sort_values(["_pk", "_base"], na_position="last")
    dropped = new[new.duplicated("_pk", keep="first")].copy()
    new = new.drop_duplicates("_pk", keep="first")
    return matched, new, dropped


def apply_fill(xlsx: Path, matched: pd.DataFrame, new: pd.DataFrame) -> dict:
    wb = load_workbook(xlsx)
    if "cln" not in wb.sheetnames:
        raise KeyError(f"{xlsx.name} 没有 cln 表（有 {wb.sheetnames}）")
    ws = wb["cln"]

    header = [c.value for c in ws[1]]
    ncol = len(header)
    pos = {}
    for c in FU_COLS:
        if c in header:
            pos[c] = header.index(c) + 1
        else:
            ncol += 1
            ws.cell(row=1, column=ncol, value=c)
            header.append(c)
            pos[c] = ncol
    col_idx = {name: i + 1 for i, name in enumerate(header)}

    # 1) 回填：按 Patient_ID
    pid_col = col_idx["Patient_ID"]
    filled, skipped = [], []
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(row=r, column=pid_col).value
        pid = "" if pid is None else str(pid).strip()
        if not pid:
            continue
        hit = matched[matched["_ini"] == pid]
        if not len(hit):
            continue
        row = hit.iloc[0]
        if row["_conf"] != "高":
            skipped.append((pid, str(row["姓名"]), row["_conf"], row["_conf_note"]))
            continue
        for c in FU_COLS:
            v = row[c]
            ws.cell(row=r, column=pos[c], value=None if _blank(v) else v)
        filled.append(pid)

    # 2) 补行
    start = ws.max_row + 1
    for k, (_, row) in enumerate(new.iterrows()):
        rr = start + k
        def put(name, val):
            if name in col_idx and not _blank(val):
                ws.cell(row=rr, column=col_idx[name], value=val)
        put("Patient_ID", row["_ini"])
        put("Gender", yt_sex(row.get("性别")))
        age = pd.to_numeric(row.get("年龄_首次"), errors="coerce")
        put("Age", None if pd.isna(age) else int(age))
        for c in FU_COLS:
            put(c, row[c])

    wb.save(xlsx)
    return {"filled": filled, "skipped": skipped, "appended": len(new),
            "cols": FU_COLS, "header": header}


def write_md(path: Path, res: dict, new: pd.DataFrame, matched: pd.DataFrame,
             dropped: pd.DataFrame, before: tuple, after: tuple) -> None:
    lines = [
        "# yt_RDL.xlsx / cln 随访列回填说明",
        "",
        f"`cln` 行数：{before[0]} → **{after[0]}**；列数：{before[1]} → **{after[1]}**",
        "",
        "## 1. 回填（按 Patient_ID = 姓名拼音首字母，且性别 + 年龄 ±3 岁双重吻合）",
        "",
        f"- 成功回填 **{len(res['filled'])}** 位：{'、'.join(sorted(res['filled']))}",
    ]
    if res["skipped"]:
        lines += ["", f"- **未回填 {len(res['skipped'])} 位**（避免张冠李戴，留空）：", ""]
        lines += ["| Patient_ID | 姓名 | 可信度 | 原因 |", "|---|---|---|---|"]
        for pid, nm, conf, note in res["skipped"]:
            lines.append(f"| {pid} | {nm} | {conf} | {note} |")
        lines += ["", "> 这些行的随访列**故意留空**。若确认是同一人，可在候选表里把"
                      "`匹配可信度` 改为「高」后重跑。"]
    lines += [
        "",
        "## 2. 新增行（cln 中原本没有的患者）",
        "",
        f"补入 **{res['appended']}** 行，追加在 `cln` 末尾。",
        "",
        "- `Patient_ID` = 姓名拼音首字母；`Gender` / `Age` 来自住院资料（Age 取最早的检查年龄）；",
        "- `Dedifferentiated` / `Recurrence` / `Reader_1` / `Reader_2` / `KI67` / `CDK4` / `MDM2` /",
        "  `Surrounding involving` / `Tumor density` / `Satellite nodules` **留空**——"
        "这些是本课题自己的判读/病理结果，不能由外部资料推断。",
        "",
        "| Patient_ID | 姓名 | 性别 | 年龄 | 复发状态 | 无复发生存时间(月) |",
        "|---|---|---|---|---|---|",
    ]
    for _, r in new.sort_values("_ini").iterrows():
        lines.append(f"| {r['_ini']} | {r['姓名']} | {r.get('性别','')} | "
                     f"{'' if _blank(r.get('年龄_首次')) else int(pd.to_numeric(r['年龄_首次']))} | "
                     f"{r['复发状态']} | {r['无复发生存时间(月)']} |")
    if len(dropped):
        lines += ["", "## 3. 一人多卡号（cln 是患者级表，只补一行）", "",
                  "| 身份证号 | 姓名 | 保留病历卡号 | 被合并掉的病历卡号 |", "|---|---|---|---|"]
        for _, r in dropped.iterrows():
            keep = new[new["_pk"] == r["_pk"]]
            lines.append(f"| {r['_pk']} | {r['姓名']} | "
                         f"{keep.iloc[0]['病历卡号'] if len(keep) else ''} | {r['病历卡号']} |")
    lines += [
        "",
        "## 4. 取值口径",
        "",
        "| 列 | 取值 |",
        "|---|---|",
        "| 无复发生存时间(月) | 复发者=基线到首次复发；未复发=基线到末次随访 |",
        "| 复发状态 | 复发 / 无复发 / 随访不足 / 无法判定 |",
        "| 复发_6月…复发_3年 | 是 / 否 / 随访不足 / 无法判定 |",
        "| `_含转移` 后缀 | 备选口径：把转移/播散/种植/进展也算事件 |",
        "",
        "## 5. ⚠ 使用注意",
        "",
        "1. **`cln` 行数变了**（118 → " + str(after[0]) + "），但 `rad` 仍是 118 行。"
        "直接重跑 RAD console 会让 Table 1 / 建模把新增的 59 行也算进去，"
        "这些行结局列全空，必须先按 `Recurrence` 非空过滤或同步更新 `rad`。",
        "2. 新增行的随访列多数是「随访不足」，**不能当作「否」**。",
        "3. 若 `无复发生存时间(月)` < 0.5，说明基线当天/次日就已有复发证据——"
        "这类患者的「基线」本身是复发就诊，做时间窗终点时应排除。",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser(description="把随访列回填进 yt_RDL.xlsx 的 cln")
    ap.add_argument("--mrn", action="store_true", help="after fill, append 病历卡号 column")
    ap.add_argument("--mrn-only", action="store_true", help="only run 病历卡号 step")
    ap.add_argument("--audit", action="store_true", help="run independent audit and exit")
    ap.add_argument("--xlsx", default=str(PATHS.YT_XLSX))
    ap.add_argument("--out", default=None, help="写到副本而不是原文件（先核对用）")
    args = ap.parse_args()
    if args.audit:
        return run_audit()
    if args.mrn_only:
        import fill_cln_mrn as _mrn
        return _mrn.main()

    target = Path(args.out) if args.out else Path(args.xlsx)
    src = Path(args.xlsx)
    if args.out:
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, target)
    else:
        PATHS.ensure_dirs()
        shutil.copy2(src, PATHS.BACKUP / src.name)
        print(f"[bak] {PATHS.BACKUP / src.name}")

    before_df = pd.read_excel(target, sheet_name="cln")
    before = before_df.shape

    fu = load_fu()
    matched, new, dropped = pick_person_rows(fu, target)
    res = apply_fill(target, matched, new)

    after_df = pd.read_excel(target, sheet_name="cln")
    after = after_df.shape

    print(f"[ok] 回填 {len(res['filled'])} 位，跳过 {len(res['skipped'])} 位")
    for s in res["skipped"]:
        print(f"     ! {s[0]} {s[1]}：{s[3]}")
    print(f"[ok] 新增 {res['appended']} 行（一人多卡号合并 {len(dropped)} 行）")
    print(f"[ok] cln: {before[0]}行×{before[1]}列 → {after[0]}行×{after[1]}列")
    print(f"[out] {target}")

    if not args.out:
        md = PATHS.REF / "yt_RDL_cln回填说明.md"
        write_md(md, res, new, matched, dropped, before, after)
        print(f"[out] {md}")
    if args.mrn:
        import fill_cln_mrn as _mrn
        print("[step] append 病历卡号…")
        rc = _mrn.main()
        if rc:
            return rc
    return 0



# ── audit (merged) ─────────────────────────────────────────────
FU_COLS = ["无复发生存时间(月)", "复发状态",
           "复发_6月", "复发_6月_含转移",
           "复发_1年", "复发_1年_含转移",
           "复发_2年", "复发_2年_含转移",
           "复发_3年", "复发_3年_含转移"]
MRN_COL = "病历卡号"
fail: list[str] = []


def chk(name, cond, detail=""):
    print(("  PASS  " if cond else "  FAIL  ") + name + (f"   {detail}" if detail else ""))
    if not cond:
        fail.append(name)


def blank(v) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan", "NaT")


def s(v) -> str:
    if blank(v):
        return ""
    t = str(v).strip()
    return t[:-2] if re.fullmatch(r"-?\d+\.0", t) else t


def run_audit() -> int:
    ap = argparse.ArgumentParser(description="核对 cln 回填结果")
    ap.add_argument("--xlsx", default=str(PATHS.YT_XLSX))
    args = ap.parse_args()

    sheets = pd.read_excel(Path(args.xlsx), sheet_name=None)
    cln = sheets["cln"]
    fu = pd.read_excel(PATHS.FU_XLSX, sheet_name="01_候选患者_合并")
    cand = pd.read_excel(PATHS.CAND_XLSX, sheet_name="01_候选患者_合并")
    for d in (fu, cand):
        d["_ini"] = d["拼音首字母"].astype(str).str.strip()

    print("### 1. 结构")
    chk("4 张表齐全", set(sheets) == {"cln", "rad", "qc", "icc"}, str(sorted(sheets)))
    chk("随访 10 列齐全", all(c in cln.columns for c in FU_COLS))
    chk(f"「{MRN_COL}」在最后一列", list(cln.columns)[-1] == MRN_COL,
        f"末列={list(cln.columns)[-1]}")

    pid = cln["Patient_ID"].astype(str).str.strip()
    print("\n### 2. 随访列")
    hi = fu[fu["匹配可信度"].astype(str).str.startswith("高")].drop_duplicates("_ini")
    hi_ids = set(hi["_ini"])
    err = 0
    for _, r in hi.iterrows():
        for _, g in cln[pid == r["_ini"]].iterrows():
            for c in FU_COLS:
                if s(r[c]) != s(g[c]):
                    err += 1
                    print(f"    {r['_ini']} {c}: 表={s(g[c])!r} 期望={s(r[c])!r}")
    chk(f"{len(hi)} 位高可信度 × 10 列逐格一致", err == 0, f"{err} 处不符")

    # 该有随访值的行 = 高可信度候选 + 当初追加进来的新患者行
    in_cohort0 = set(cand.loc[cand["匹配yt_RDL_Patient_ID"].notna() &
                              (cand["匹配yt_RDL_Patient_ID"].astype(str).str.strip() != ""), "_ini"])
    appended0 = set(cand.loc[~cand["_ini"].isin(in_cohort0), "_ini"])
    expect_filled = hi_ids | appended0
    other = cln[~pid.isin(expect_filled)]
    has_v = other[FU_COLS].map(lambda v: not blank(v)).any(axis=1)
    chk("既非高可信度、也非追加新患者的行，随访列全空", not has_v.any(),
        f"{int(has_v.sum())} 行有值")
    sel = cln[pid.isin(expect_filled)]
    has2 = sel[FU_COLS].map(lambda v: not blank(v)).any(axis=1)
    chk("该有随访值的行都有值", bool(has2.all()),
        f"{int((~has2).sum())} 行缺")

    print("\n### 3. 病历卡号列")
    got = cln[MRN_COL].map(s)
    in_cohort = set(cand.loc[cand["匹配yt_RDL_Patient_ID"].notna() &
                            (cand["匹配yt_RDL_Patient_ID"].astype(str).str.strip() != ""), "_ini"])
    all_ini = set(cand["_ini"])
    should = pid.map(lambda p: p in hi_ids or (p not in in_cohort and p in all_ini))
    bad = [i for i in range(len(cln)) if should.iloc[i] and not got.iloc[i]]
    extra = [i for i in range(len(cln)) if (not should.iloc[i]) and got.iloc[i]]
    chk("该填的行都填了", not bad, f"{len(bad)} 行缺，样例 {[pid.iloc[i] for i in bad[:5]]}")
    chk("不该填的行都空着", not extra, f"{len(extra)} 行多填，样例 {[pid.iloc[i] for i in extra[:5]]}")
    cand_mrn = set(cand[MRN_COL].map(s))
    unknown = sorted(set(got) - {""} - cand_mrn)
    chk("填进去的病历卡号都能在候选表找到", not unknown, f"未知 {unknown[:5]}")
    print(f"  已填 {int(got.ne('').sum())}/{len(cln)} 行，唯一值 {got[got.ne('')].nunique()} 个")

    print("\n### 4. 其他表")
    for t, shp in [("rad", (118, 511)), ("qc", (122, 37)), ("icc", (1316, 9))]:
        chk(f"{t} 结构完好", sheets[t].shape == shp, f"{sheets[t].shape} vs {shp}")

    print("\n" + "=" * 70)
    print("FAILED:", fail if fail else "无 —— 全部检查通过")
    return 1 if fail else 0

if __name__ == "__main__":
    raise SystemExit(main())
